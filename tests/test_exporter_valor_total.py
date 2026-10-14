"""VALOR_TOTAL nunca sai vazio quando quantidade e preço existem.

O importador de Excel do Fire converte a coluna VALOR_TOTAL em número; célula vazia
estoura `'' is not a valid floating point value` e o pedido entra pela metade (OC Beira
Rio 14330456, 02/10/2026: 1 item gravado, 2 tentativas). Todo parser que lê o total da
linha do documento (Riachuelo, SBF, Sam's, Nasmar...) já saía preenchido; a Beira Rio,
cuja OC só imprime o total do pedido, saía vazio.
"""

from __future__ import annotations

import tempfile
from pathlib import Path

import openpyxl
import pytest

from app.exporters.erp_exporter import HEADERS, ERPExporter
from app.models.order import Order, OrderHeader, OrderItem

SAMPLES = Path(__file__).parent.parent / "samples"


def _export(items: list[OrderItem]) -> list[tuple]:
    order = Order(
        header=OrderHeader(
            order_number="T-1",
            customer_name="Cliente Teste",
            customer_cnpj="88.379.771/0041-70",
        ),
        items=items,
    )
    with tempfile.TemporaryDirectory() as tmp:
        files = ERPExporter().export(order, tmp)
        ws = openpyxl.load_workbook(files[0]).active
        col = [c.value for c in ws[1]].index("VALOR_TOTAL")
        return [r[col] for r in ws.iter_rows(min_row=2, values_only=True)]


def _item(qty, price, total=None) -> OrderItem:
    return OrderItem(
        product_code="P1", description="d", quantity=qty, unit_price=price, total_price=total
    )


def test_total_ausente_vira_quantidade_vezes_preco():
    assert _export([_item(3300, 8.59)]) == [28347.0]


def test_total_calculado_e_arredondado_em_centavos():
    # 3 * 0.1 em float é 0.30000000000000004
    assert _export([_item(3, 0.1)]) == [0.3]


def test_total_que_veio_do_documento_e_preservado():
    # o documento manda: não recalcula por cima (ex.: desconto ou arredondamento do cliente)
    assert _export([_item(10, 2.0, total=19.5)]) == [19.5]


def test_sem_preco_nao_inventa_total():
    assert _export([_item(10, None)]) == [None]


def test_coluna_valor_total_existe():
    assert "VALOR_TOTAL" in HEADERS


def test_beira_rio_grade_unica_nenhuma_linha_sem_total():
    from app.ingestion.file_loader import LoadedFile
    from app.pipeline import process

    p = SAMPLES / "PEDIDO BEIRA RIO 14330456.pdf"
    if not p.exists():
        pytest.skip("sample ausente")
    order = process(LoadedFile(path=p, extension=".pdf", raw=p.read_bytes()))
    with tempfile.TemporaryDirectory() as tmp:
        files = ERPExporter().export(order, tmp)
        ws = openpyxl.load_workbook(files[0]).active
        col = [c.value for c in ws[1]].index("VALOR_TOTAL")
        totais = [r[col] for r in ws.iter_rows(min_row=2, values_only=True)]
    assert len(totais) == 16
    assert all(t is not None for t in totais)
    # soma das linhas = `Total: 445.050,00` impresso no rodapé da OC
    assert round(sum(totais), 2) == 445050.00
