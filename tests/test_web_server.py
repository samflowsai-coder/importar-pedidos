"""Integration tests for the FastAPI web server."""

from __future__ import annotations

from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from app.persistence import db
from app.web.server import app

client = TestClient(app, raise_server_exceptions=False)

SAMPLES = Path(__file__).parent.parent / "samples"


@pytest.fixture(autouse=True)
def isolated_sqlite(tmp_path_factory):
    """Redirect app_state.db to a per-test tmp file so tests never touch prod DB."""
    tmp = tmp_path_factory.mktemp("dbstate")
    db.set_db_path(tmp / "app_state.db")
    db.reset_init_cache()
    yield
    db.set_db_path(None)
    db.reset_init_cache()


# ── Basic smoke ──────────────────────────────────────────────────────────────


def test_index_returns_html():
    r = client.get("/")
    assert r.status_code == 200
    assert "text/html" in r.headers["content-type"]
    assert "Portal de Pedidos" in r.text


def test_env_edit_page_has_flowpcp_section():
    r = client.get("/admin/ambientes/any-id")
    assert r.status_code == 200
    assert "Integração FlowPCP" in r.text
    assert 'name="flowpcp_enabled"' in r.text
    assert 'name="flowpcp_service_token"' in r.text


def test_config_returns_default_output_dir():
    r = client.get("/api/config")
    assert r.status_code == 200
    data = r.json()
    assert "outputDir" in data
    assert "output" in data["outputDir"]
    assert "exportMode" in data
    assert data["exportMode"] in ("xlsx", "db", "both")


# ── Security: download endpoint ──────────────────────────────────────────────
# Achado pré-existente (não criado por esta entrega): a rota resolvia
# qualquer caminho absoluto e só checava o sufixo `.xlsx` — sem
# `require_user` e sem confinamento de raiz. Corrigido no mesmo padrão que
# `baixar_arquivo_original` já usa pra `recebidos/`: raízes = `OUTPUT_DIR`
# legado + `output_dir` de cada ambiente ATIVO (as duas fontes de onde
# `ERPExporter().export(...)` grava — ver `_run_exporters` e `_send_one_to_fire`).


def test_download_rejects_non_xlsx():
    r = client.get("/api/download?path=/etc/passwd")
    assert r.status_code == 403


def test_download_rejects_paths_outside_allowed_roots():
    """Sufixo certo não basta: fora de `OUTPUT_DIR`/`output_dir` dos
    ambientes é 403, mesmo sem o arquivo existir — a checagem de raiz roda
    ANTES da checagem de existência (mesma ordem de `baixar_arquivo_original`),
    pra não vazar "existe/não existe" de fora das pastas de saída."""
    r = client.get("/api/download?path=/etc/hosts.xlsx")
    assert r.status_code == 403


def test_download_missing_xlsx_returns_404(tmp_path, monkeypatch):
    """Dentro da raiz permitida, mas ausente no disco: 404, não 403."""
    from app import config as app_config

    monkeypatch.setattr(app_config, "_CONFIG_FILE", tmp_path / "config.json")
    monkeypatch.setenv("OUTPUT_DIR", str(tmp_path))

    r = client.get(f"/api/download?path={tmp_path / 'missing.xlsx'}")
    assert r.status_code == 404


def test_download_allows_xlsx_inside_legacy_output_dir(tmp_path, monkeypatch):
    """Autenticado + caminho dentro do `OUTPUT_DIR` legado: 200."""
    from app import config as app_config

    monkeypatch.setattr(app_config, "_CONFIG_FILE", tmp_path / "config.json")
    monkeypatch.setenv("OUTPUT_DIR", str(tmp_path))

    arquivo = tmp_path / "PEDIDO-123.xlsx"
    arquivo.write_bytes(b"fake-xlsx-bytes")

    r = client.get(f"/api/download?path={arquivo}")
    assert r.status_code == 200
    assert r.content == b"fake-xlsx-bytes"
    assert "spreadsheetml" in r.headers["content-type"]


def test_download_allows_xlsx_inside_environment_output_dir(tmp_path, monkeypatch):
    """Confinamento também aceita o `output_dir` de um ambiente ATIVO, não só
    o `OUTPUT_DIR` legado — `_run_exporters`/`_send_one_to_fire` gravam ali
    quando a request tem ambiente selecionado (multi-empresa)."""
    from app.persistence import environments_repo, router

    monkeypatch.setenv("APP_DATA_DIR", str(tmp_path))
    router.reset_init_cache()
    with router.shared_connect():
        pass

    env_out = tmp_path / "env-out"
    env_out.mkdir()
    environments_repo.create(
        slug="download-test-env",
        name="Download Test Env",
        watch_dir=str(tmp_path / "env-watch"),
        output_dir=str(env_out),
        fb_path="",
    )

    arquivo = env_out / "PEDIDO-999.xlsx"
    arquivo.write_bytes(b"fake-xlsx-env")

    r = client.get(f"/api/download?path={arquivo}")
    assert r.status_code == 200
    assert r.content == b"fake-xlsx-env"


# ── Security: upload validation ──────────────────────────────────────────────


def test_upload_rejects_disallowed_extension(tmp_path):
    r = client.post(
        "/api/process",
        data={"output_dir": str(tmp_path)},
        files=[("files", ("malware.exe", b"MZ payload", "application/octet-stream"))],
    )
    assert r.status_code == 200
    data = r.json()
    assert data["results"] == []
    assert any("suportado" in e["error"].lower() for e in data["errors"])


def test_upload_rejects_oversized_file(tmp_path):
    big = b"0" * (51 * 1024 * 1024)  # 51 MB
    r = client.post(
        "/api/process",
        data={"output_dir": str(tmp_path)},
        files=[("files", ("big.pdf", big, "application/pdf"))],
    )
    assert r.status_code == 200
    data = r.json()
    assert any("limite" in e["error"].lower() for e in data["errors"])


# ── Filesystem browser ───────────────────────────────────────────────────────


def test_fs_returns_directories():
    r = client.get("/api/fs?path=~")
    assert r.status_code == 200
    data = r.json()
    assert "current" in data
    assert isinstance(data["entries"], list)
    # Entries must all be directories (names only, no file content exposed)
    for entry in data["entries"]:
        assert "name" in entry
        assert "path" in entry


def test_fs_handles_nonexistent_path():
    r = client.get("/api/fs?path=/this/does/not/exist/ever")
    assert r.status_code == 200  # falls back to parent


def test_fs_requires_admin_auth_without_bypass(real_auth):
    c = TestClient(app)
    r = c.get("/api/fs?path=~&file_ext=.fdb")
    assert r.status_code == 401


# ── End-to-end: real PDFs ────────────────────────────────────────────────────


@pytest.mark.skipif(not SAMPLES.exists(), reason="samples/ directory not found")
def test_process_calcenter_pdf(tmp_path):
    pdf = SAMPLES / "2600009562-2026-02-25.pdf"
    if not pdf.exists():
        pytest.skip("Sample PDF not available")

    r = client.post(
        "/api/process",
        data={"output_dir": str(tmp_path)},
        files=[("files", (pdf.name, pdf.read_bytes(), "application/pdf"))],
    )
    assert r.status_code == 200
    data = r.json()
    assert len(data["results"]) == 1
    result = data["results"][0]
    assert result["order"] == "2600009562"
    assert len(result["files"]) == 1
    assert result["files"][0]["name"].endswith(".xlsx")
    assert (tmp_path / result["files"][0]["name"]).exists()


@pytest.mark.skipif(not SAMPLES.exists(), reason="samples/ directory not found")
def test_process_riachuelo_splits_into_three(tmp_path):
    pdf = SAMPLES / "PEDIDO 6702604130.pdf"
    if not pdf.exists():
        pytest.skip("Sample PDF not available")

    r = client.post(
        "/api/process",
        data={"output_dir": str(tmp_path)},
        files=[("files", (pdf.name, pdf.read_bytes(), "application/pdf"))],
    )
    assert r.status_code == 200
    data = r.json()
    assert len(data["results"]) == 1
    assert len(data["results"][0]["files"]) == 3


# ── Preview → Commit flow ────────────────────────────────────────────────────


@pytest.mark.skipif(not SAMPLES.exists(), reason="samples/ directory not found")
def test_preview_returns_structured_payload():
    pdf = SAMPLES / "2600009562-2026-02-25.pdf"
    if not pdf.exists():
        pytest.skip("Sample PDF not available")

    r = client.post(
        "/api/preview",
        files=[("file", (pdf.name, pdf.read_bytes(), "application/pdf"))],
    )
    assert r.status_code == 200
    data = r.json()
    assert "preview_id" in data
    assert data["header"]["order_number"] == "2600009562"
    assert isinstance(data["items"], list)
    assert data["totals"]["items_count"] == len(data["items"])
    assert "groups" in data


def test_preview_rejects_disallowed_extension():
    r = client.post(
        "/api/preview",
        files=[("file", ("malware.exe", b"MZ", "application/octet-stream"))],
    )
    assert r.status_code == 400


def test_commit_rejects_unknown_preview_id():
    r = client.post("/api/commit", json={"preview_id": "does-not-exist"})
    assert r.status_code == 404


@pytest.mark.skipif(not SAMPLES.exists(), reason="samples/ directory not found")
def test_commit_consumes_preview_and_persists_log(tmp_path, monkeypatch):
    from app import config as app_config

    # Redirect config to tmp so we don't touch real log / output paths
    monkeypatch.setenv("OUTPUT_DIR", str(tmp_path))
    monkeypatch.setenv("EXPORT_MODE", "xlsx")
    monkeypatch.setattr(
        app_config,
        "load",
        lambda: {
            "watch_dir": str(tmp_path),
            "output_dir": str(tmp_path),
            "export_mode": "xlsx",
        },
    )
    monkeypatch.setattr(
        app_config,
        "imported_dir",
        lambda _cfg: tmp_path,
    )

    pdf = SAMPLES / "2600009562-2026-02-25.pdf"
    if not pdf.exists():
        pytest.skip("Sample PDF not available")

    r = client.post(
        "/api/preview",
        files=[("file", (pdf.name, pdf.read_bytes(), "application/pdf"))],
    )
    assert r.status_code == 200
    preview_id = r.json()["preview_id"]

    r2 = client.post("/api/commit", json={"preview_id": preview_id})
    assert r2.status_code == 200, r2.text
    assert r2.json()["order"] == "2600009562"
    assert r2.json()["portal_status"] == "parsed"

    # Second commit with same id must be rejected
    r3 = client.post("/api/commit", json={"preview_id": preview_id})
    assert r3.status_code == 409

    # Entry is persisted and queryable via /api/imported, with portal_status='parsed'
    r4 = client.get("/api/imported?limit=10")
    assert r4.status_code == 200
    body = r4.json()
    assert body["total"] >= 1
    found = [e for e in body["entries"] if e["order_number"] == "2600009562"]
    assert found
    assert found[0]["portal_status"] == "parsed"
    assert found[0]["fire_codigo"] is None


@pytest.mark.skipif(not SAMPLES.exists(), reason="samples/ directory not found")
def test_preview_pending_reads_from_watch_folder(tmp_path, monkeypatch):
    from app import config as app_config

    # Seed watch folder with a sample PDF
    sample = SAMPLES / "2600009562-2026-02-25.pdf"
    if not sample.exists():
        pytest.skip("Sample PDF not available")
    watch = tmp_path / "watch"
    watch.mkdir()
    imp = tmp_path / "imported"
    imp.mkdir()
    seeded = watch / sample.name
    seeded.write_bytes(sample.read_bytes())

    monkeypatch.setattr(
        app_config,
        "load",
        lambda: {
            "watch_dir": str(watch),
            "output_dir": str(tmp_path),
            "export_mode": "xlsx",
        },
    )
    monkeypatch.setattr(app_config, "imported_dir", lambda _cfg: imp)

    r = client.post("/api/preview-pending", json={"filename": sample.name})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["header"]["order_number"] == "2600009562"
    preview_id = body["preview_id"]

    r2 = client.post("/api/commit", json={"preview_id": preview_id, "outputDir": str(tmp_path)})
    assert r2.status_code == 200, r2.text

    # File must have been moved from watch to imported folder
    assert not seeded.exists()
    assert (imp / sample.name).exists()


@pytest.mark.skipif(not SAMPLES.exists(), reason="samples/ directory not found")
def test_preview_pending_uses_selected_environment_watch_dir(tmp_path, monkeypatch):
    from app.persistence import environments_repo

    sample = SAMPLES / "2600009562-2026-02-25.pdf"
    if not sample.exists():
        pytest.skip("Sample PDF not available")

    legacy_watch = tmp_path / "legacy-watch"
    env_watch = tmp_path / "env-watch"
    legacy_watch.mkdir()
    env_watch.mkdir()
    seeded = env_watch / sample.name
    seeded.write_bytes(sample.read_bytes())

    env = environments_repo.create(
        slug="mm",
        name="MM",
        watch_dir=str(env_watch),
        output_dir=str(tmp_path / "env-out"),
        fb_path="",
    )
    monkeypatch.setattr(
        "app.config.load",
        lambda: {
            "watch_dir": str(legacy_watch),
            "output_dir": str(tmp_path / "legacy-out"),
            "export_mode": "xlsx",
        },
    )

    client.cookies.set("portal_env", env["id"])
    try:
        r = client.post("/api/preview-pending", json={"filename": sample.name})
    finally:
        client.cookies.clear()
    assert r.status_code == 200, r.text
    assert r.json()["header"]["order_number"] == "2600009562"


def test_config_reflete_firebird_do_ambiente_ativo(tmp_path):
    from app.persistence import environments_repo

    env = environments_repo.create(
        slug="mmfire", name="MM Americanense",
        watch_dir=str(tmp_path / "w"), output_dir=str(tmp_path / "o"),
        fb_path="/srv/mm.fdb",
    )
    client.cookies.set("portal_env", env["id"])
    try:
        r = client.get("/api/config")
    finally:
        client.cookies.clear()
    assert r.status_code == 200
    body = r.json()
    assert body["firebirdConfigured"] is True
    assert body["environment"]["name"] == "MM Americanense"
    assert body["environment"]["slug"] == "mmfire"
    assert body["watchDir"] == str(tmp_path / "w")


def test_config_ambiente_sem_fire_reporta_sem_banco(tmp_path):
    from app.persistence import environments_repo

    env = environments_repo.create(
        slug="nofire", name="Sem Fire",
        watch_dir=str(tmp_path / "w"), output_dir=str(tmp_path / "o"),
        fb_path="",
    )
    client.cookies.set("portal_env", env["id"])
    try:
        r = client.get("/api/config")
    finally:
        client.cookies.clear()
    assert r.status_code == 200
    body = r.json()
    assert body["firebirdConfigured"] is False
    assert body["environment"]["name"] == "Sem Fire"


def test_html_pages_sao_no_cache():
    # Sem isso o browser servia o HTML antigo pós-deploy (mudança de UI só com
    # hard-reload). no-cache força revalidação → o HTML novo é pego sozinho.
    r = client.get("/")
    assert r.status_code == 200
    assert "text/html" in r.headers.get("content-type", "")
    assert r.headers.get("cache-control") == "no-cache"


def test_static_assets_nao_sao_no_cache():
    # Os assets seguem cacheáveis — o cache-bust deles é o ?v=N no HTML.
    r = client.get("/static/js/shell.js")
    assert r.status_code == 200
    assert r.headers.get("cache-control") != "no-cache"


def test_preview_pending_rejects_missing_file(tmp_path, monkeypatch):
    from app import config as app_config

    monkeypatch.setattr(
        app_config,
        "load",
        lambda: {
            "watch_dir": str(tmp_path),
            "output_dir": str(tmp_path),
            "export_mode": "xlsx",
        },
    )
    r = client.post("/api/preview-pending", json={"filename": "nao-existe.pdf"})
    assert r.status_code == 404


def test_preview_pending_rejects_path_traversal(tmp_path, monkeypatch):
    from app import config as app_config

    monkeypatch.setattr(
        app_config,
        "load",
        lambda: {
            "watch_dir": str(tmp_path),
            "output_dir": str(tmp_path),
            "export_mode": "xlsx",
        },
    )
    # Attempt to escape the watch folder
    r = client.post("/api/preview-pending", json={"filename": "../../../etc/passwd"})
    assert r.status_code in (400, 404)  # either ext-reject or not-found after basename strip


def test_cancel_parsed_order_marks_as_cancelled():
    import uuid
    from datetime import datetime

    from app.persistence import repo

    entry_id = str(uuid.uuid4())
    repo.insert_import(
        {
            "id": entry_id,
            "source_filename": "x.pdf",
            "imported_at": datetime.now().isoformat(timespec="seconds"),
            "order_number": "TEST-1",
            "customer": "ACME",
            "status": "success",
            "portal_status": "parsed",
            "snapshot": {"header": {"order_number": "TEST-1"}, "items": []},
        }
    )

    r = client.post(f"/api/imported/{entry_id}/cancel", json={"reason": "duplicado"})
    assert r.status_code == 200
    assert r.json()["portal_status"] == "cancelled"

    got = repo.get_import(entry_id)
    assert got["portal_status"] == "cancelled"


def test_cancel_sent_to_fire_rejected():
    import uuid
    from datetime import datetime

    from app.persistence import repo

    entry_id = str(uuid.uuid4())
    repo.insert_import(
        {
            "id": entry_id,
            "source_filename": "x.pdf",
            "imported_at": datetime.now().isoformat(timespec="seconds"),
            "order_number": "TEST-2",
            "status": "success",
            "portal_status": "sent_to_fire",
            "fire_codigo": 999,
            "snapshot": {"header": {"order_number": "TEST-2"}, "items": []},
        }
    )
    r = client.post(f"/api/imported/{entry_id}/cancel", json={})
    assert r.status_code == 409


def test_cancel_found_in_fire_rejected_without_audit():
    """found_in_fire (reconciliado, não inserido pelo portal) não pode ser
    cancelado — e a tentativa recusada não deve gravar audit 'cancelled',
    porque o cancelamento não aconteceu."""
    import uuid
    from datetime import datetime

    from app.persistence import repo

    entry_id = str(uuid.uuid4())
    repo.insert_import(
        {
            "id": entry_id,
            "source_filename": "x.pdf",
            "imported_at": datetime.now().isoformat(timespec="seconds"),
            "order_number": "TEST-FOUND",
            "status": "success",
            "portal_status": "found_in_fire",
            "fire_codigo": 555,
            "snapshot": {"header": {"order_number": "TEST-FOUND"}, "items": []},
        }
    )
    r = client.post(f"/api/imported/{entry_id}/cancel", json={"reason": "duplicado"})
    assert r.status_code == 409

    got = repo.get_import(entry_id)
    assert got["portal_status"] == "found_in_fire"

    audit_events = [a["event_type"] for a in repo.list_audit(entry_id)]
    assert "cancelled" not in audit_events


def test_send_to_fire_rejects_wrong_portal_status():
    import uuid
    from datetime import datetime

    from app.persistence import repo

    entry_id = str(uuid.uuid4())
    repo.insert_import(
        {
            "id": entry_id,
            "source_filename": "x.pdf",
            "imported_at": datetime.now().isoformat(timespec="seconds"),
            "order_number": "TEST-3",
            "status": "success",
            "portal_status": "cancelled",
            "snapshot": {"header": {"order_number": "TEST-3"}, "items": []},
        }
    )
    r = client.post(f"/api/imported/{entry_id}/send-to-fire")
    assert r.status_code == 409


def test_send_to_fire_missing_order_returns_404():
    r = client.post("/api/imported/does-not-exist/send-to-fire")
    assert r.status_code == 404


def test_send_to_fire_inserts_when_success(monkeypatch):
    """Mock FirebirdExporter to simulate a real Fire insert without the DB."""
    import uuid
    from datetime import datetime

    from app.exporters import firebird_exporter as fb_mod
    from app.persistence import repo

    entry_id = str(uuid.uuid4())
    repo.insert_import(
        {
            "id": entry_id,
            "source_filename": "pedido.pdf",
            "imported_at": datetime.now().isoformat(timespec="seconds"),
            "order_number": "TEST-OK",
            "customer": "ACME",
            "status": "success",
            "portal_status": "parsed",
            "snapshot": {
                "header": {"order_number": "TEST-OK", "customer_name": "ACME"},
                "items": [{"description": "x", "quantity": 1.0, "ean": "1234"}],
                "source_file": "",
            },
        }
    )

    def _fake_export(self, order, *, override_client_id=None, produtos_por_item=None):
        return fb_mod.FirebirdExportResult(
            order_number=order.header.order_number,
            items_inserted=1,
            fire_codigo=4242,
        )

    monkeypatch.setattr(fb_mod.FirebirdExporter, "export", _fake_export)

    # also force export_mode to 'db' to skip XLSX export path
    from app import config as app_config

    monkeypatch.setattr(
        app_config,
        "load",
        lambda: {
            "watch_dir": ".",
            "output_dir": ".",
            "export_mode": "db",
        },
    )

    r = client.post(f"/api/imported/{entry_id}/send-to-fire")
    assert r.status_code == 200, r.text
    assert r.json()["fire_codigo"] == 4242

    got = repo.get_import(entry_id)
    assert got["portal_status"] == "sent_to_fire"
    assert got["fire_codigo"] == 4242


def _seed_parsed_import(repo, entry_id: str) -> None:
    from datetime import datetime

    repo.insert_import(
        {
            "id": entry_id,
            "source_filename": "pedido.pdf",
            "imported_at": datetime.now().isoformat(timespec="seconds"),
            "order_number": "TEST-FLOW",
            "customer": "MM",
            "status": "success",
            "portal_status": "parsed",
            "snapshot": {
                "header": {"order_number": "TEST-FLOW", "customer_name": "MM"},
                "items": [{"description": "x", "quantity": 1.0, "ean": "1234"}],
                "source_file": "",
            },
        }
    )


def _stub_fire_success(monkeypatch):
    from app.exporters import firebird_exporter as fb_mod

    monkeypatch.setattr(
        fb_mod.FirebirdExporter,
        "export",
        lambda self, order, *, override_client_id=None, produtos_por_item=None: fb_mod.FirebirdExportResult(
            order_number=order.header.order_number, items_inserted=1, fire_codigo=4242
        ),
    )
    # Neutraliza price-check (sem Fire real) — deterministico.
    monkeypatch.setattr("app.erp.product_check.check_order", lambda order, env=None: {})
    monkeypatch.setattr(
        "app.erp.product_check.is_blocking", lambda check, ack_items=None: (False, {})
    )


def test_send_to_fire_pushes_to_flowpcp_for_env_with_slug(monkeypatch, tmp_path):
    import uuid
    from unittest.mock import MagicMock

    from app.persistence import repo
    from app.web import server

    entry_id = str(uuid.uuid4())
    _seed_parsed_import(repo, entry_id)
    _stub_fire_success(monkeypatch)

    spy = MagicMock()
    monkeypatch.setattr(server, "push_new_order", spy)

    cfg = {"output_dir": str(tmp_path), "export_mode": "db"}
    outcome = server._send_one_to_fire(entry_id, cfg, request_env={"id": "e1", "slug": "mm"})

    assert outcome.ok
    spy.assert_called_once()
    args, kwargs = spy.call_args
    assert kwargs["import_id"] == entry_id
    assert kwargs["slug"] == "mm"
    assert args[0].header.order_number == "TEST-FLOW"  # o Order


def test_send_to_fire_skips_flowpcp_without_env(monkeypatch, tmp_path):
    import uuid
    from unittest.mock import MagicMock

    from app.persistence import repo
    from app.web import server

    entry_id = str(uuid.uuid4())
    _seed_parsed_import(repo, entry_id)
    _stub_fire_success(monkeypatch)

    spy = MagicMock()
    monkeypatch.setattr(server, "push_new_order", spy)

    cfg = {"output_dir": str(tmp_path), "export_mode": "db"}
    outcome = server._send_one_to_fire(entry_id, cfg, request_env=None)

    assert outcome.ok
    spy.assert_not_called()


def test_batch_send_to_fire_mixed_outcomes(monkeypatch):
    """Batch endpoint tolerates partial failures: some parsed, one cancelled, one not-found."""
    import uuid
    from datetime import datetime

    from app import config as app_config
    from app.exporters import firebird_exporter as fb_mod
    from app.persistence import repo

    parsed_ids = []
    for _ in range(2):
        entry_id = str(uuid.uuid4())
        repo.insert_import(
            {
                "id": entry_id,
                "source_filename": "pedido.pdf",
                "imported_at": datetime.now().isoformat(timespec="seconds"),
                "order_number": f"BATCH-{entry_id[:4]}",
                "customer": "ACME",
                "status": "success",
                "portal_status": "parsed",
                "snapshot": {
                    "header": {"order_number": f"BATCH-{entry_id[:4]}", "customer_name": "ACME"},
                    "items": [{"description": "x", "quantity": 1.0}],
                    "source_file": "",
                },
            }
        )
        parsed_ids.append(entry_id)

    cancelled_id = str(uuid.uuid4())
    repo.insert_import(
        {
            "id": cancelled_id,
            "source_filename": "x.pdf",
            "imported_at": datetime.now().isoformat(timespec="seconds"),
            "order_number": "CANC-1",
            "status": "success",
            "portal_status": "cancelled",
            "snapshot": {"header": {"order_number": "CANC-1"}, "items": []},
        }
    )

    fire_seq = iter([4001, 4002])

    def _fake_export(self, order, *, override_client_id=None, produtos_por_item=None):
        return fb_mod.FirebirdExportResult(
            order_number=order.header.order_number,
            items_inserted=1,
            fire_codigo=next(fire_seq),
        )

    monkeypatch.setattr(fb_mod.FirebirdExporter, "export", _fake_export)
    monkeypatch.setattr(
        app_config,
        "load",
        lambda: {
            "watch_dir": ".",
            "output_dir": ".",
            "export_mode": "db",
        },
    )

    r = client.post(
        "/api/batch/send-to-fire",
        json={
            "ids": parsed_ids + [cancelled_id, "unknown-id"],
        },
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["total"] == 4
    assert body["ok"] == 2
    assert body["failed"] == 2

    by_id = {r["id"]: r for r in body["results"]}
    for pid in parsed_ids:
        assert by_id[pid]["ok"] is True
        assert by_id[pid]["fire_codigo"] in (4001, 4002)
    assert by_id[cancelled_id]["ok"] is False
    assert by_id[cancelled_id]["reason"] == "wrong_status"
    assert by_id["unknown-id"]["ok"] is False
    assert by_id["unknown-id"]["reason"] == "not_found"

    # Parsed rows are now sent_to_fire
    for pid in parsed_ids:
        got = repo.get_import(pid)
        assert got["portal_status"] == "sent_to_fire"


def test_export_xlsx_generates_files_without_calling_firebird(monkeypatch, tmp_path):
    """xlsx-only flow: ERPExporter runs, FirebirdExporter is NEVER instantiated/called."""
    import uuid
    from datetime import datetime

    from app import config as app_config
    from app.exporters import erp_exporter as erp_mod
    from app.exporters import firebird_exporter as fb_mod
    from app.persistence import repo

    entry_id = str(uuid.uuid4())
    repo.insert_import(
        {
            "id": entry_id,
            "source_filename": "pedido.pdf",
            "imported_at": datetime.now().isoformat(timespec="seconds"),
            "order_number": "XLSX-OK",
            "customer": "ACME",
            "status": "success",
            "portal_status": "parsed",
            "snapshot": {
                "header": {"order_number": "XLSX-OK", "customer_name": "ACME"},
                "items": [{"description": "x", "quantity": 1.0}],
                "source_file": "",
            },
        }
    )

    # Stub ERPExporter to return predictable file paths without touching disk.
    erp_calls: list[str] = []

    def _fake_erp_export(self, order, output_dir):
        erp_calls.append(order.header.order_number or "?")
        out = Path(output_dir) / "ACME_XLSX-OK.xlsx"
        out.parent.mkdir(parents=True, exist_ok=True)
        out.touch()
        return [out]

    monkeypatch.setattr(erp_mod.ERPExporter, "export", _fake_erp_export)

    # Hard-fail if Firebird is touched.
    def _explode(self, order, *, override_client_id=None, produtos_por_item=None):
        raise AssertionError("FirebirdExporter.export must not be called in xlsx mode")

    monkeypatch.setattr(fb_mod.FirebirdExporter, "export", _explode)

    monkeypatch.setattr(
        app_config,
        "load",
        lambda: {
            "watch_dir": str(tmp_path),
            "output_dir": str(tmp_path),
            "export_mode": "xlsx",
        },
    )

    r = client.post(f"/api/imported/{entry_id}/export-xlsx")
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["entry_id"] == entry_id
    assert body["portal_status"] == "parsed"
    assert body["output_files"] and body["output_files"][0]["name"].endswith(".xlsx")
    assert erp_calls == ["XLSX-OK"]

    # State unchanged: still 'parsed', no fire_codigo set.
    got = repo.get_import(entry_id)
    assert got["portal_status"] == "parsed"
    assert got.get("fire_codigo") in (None, 0)


def test_export_xlsx_rejects_wrong_portal_status():
    import uuid
    from datetime import datetime

    from app.persistence import repo

    entry_id = str(uuid.uuid4())
    repo.insert_import(
        {
            "id": entry_id,
            "source_filename": "x.pdf",
            "imported_at": datetime.now().isoformat(timespec="seconds"),
            "order_number": "XLSX-WRONG",
            "status": "success",
            "portal_status": "cancelled",
            "snapshot": {"header": {"order_number": "XLSX-WRONG"}, "items": []},
        }
    )
    r = client.post(f"/api/imported/{entry_id}/export-xlsx")
    assert r.status_code == 409


def test_export_xlsx_rejects_found_in_fire_with_explanation():
    """found_in_fire já está no ERP — reexportar convida duplicata. A mensagem
    precisa explicar o motivo (não é erro do usuário)."""
    import uuid
    from datetime import datetime

    from app.persistence import repo

    entry_id = str(uuid.uuid4())
    repo.insert_import(
        {
            "id": entry_id,
            "source_filename": "x.pdf",
            "imported_at": datetime.now().isoformat(timespec="seconds"),
            "order_number": "XLSX-FOUND",
            "status": "success",
            "portal_status": "found_in_fire",
            "fire_codigo": 777,
            "snapshot": {"header": {"order_number": "XLSX-FOUND"}, "items": []},
        }
    )
    r = client.post(f"/api/imported/{entry_id}/export-xlsx")
    assert r.status_code == 409
    detail = r.json()["detail"].lower()
    # Mensagem dedicada, não a genérica de "wrong_status": precisa explicar
    # que o pedido já está no Fire (não é erro do usuário reexportar).
    assert "já" in detail and "fire" in detail
    assert "em revisão" not in detail


def test_batch_send_to_fire_rejects_empty_and_oversized():
    r = client.post("/api/batch/send-to-fire", json={"ids": []})
    assert r.status_code == 400

    r2 = client.post("/api/batch/send-to-fire", json={"ids": ["x"] * 101})
    assert r2.status_code == 400


def test_rehydrate_preview_returns_snapshot():
    import uuid
    from datetime import datetime

    from app.persistence import repo

    entry_id = str(uuid.uuid4())
    repo.insert_import(
        {
            "id": entry_id,
            "source_filename": "x.pdf",
            "imported_at": datetime.now().isoformat(timespec="seconds"),
            "order_number": "REHYDR-1",
            "customer": "ACME",
            "status": "success",
            "portal_status": "parsed",
            "snapshot": {
                "header": {"order_number": "REHYDR-1", "customer_name": "ACME"},
                "items": [{"description": "item A", "quantity": 3.0}],
                "source_file": "",
            },
            "check": {
                "available": False,
                "reason": "FB_DATABASE_NOT_SET",
                "client": {"match": False},
                "items": [],
                "summary": {},
            },
        }
    )
    r = client.get(f"/api/imported/{entry_id}/preview")
    assert r.status_code == 200
    body = r.json()
    assert body["preview_id"] == entry_id
    assert body["portal_status"] == "parsed"
    assert body["header"]["order_number"] == "REHYDR-1"
    assert body["check"]["available"] is False


def test_rehydrate_preview_includes_fire_status_last_seen():
    """Um pedido `found_in_fire` CANCELADO no Fire reconcilia e some da fila de
    revisão — corretamente, porque existe lá. Mas se o modal de detalhe não
    mostrar o status, a operadora conclui, errado, que está tudo resolvido.
    O selo da lista já mostrava `fire_status_last_seen`; o modal precisa do
    mesmo dado no payload."""
    import uuid
    from datetime import datetime

    from app.persistence import repo

    entry_id = str(uuid.uuid4())
    repo.insert_import(
        {
            "id": entry_id,
            "source_filename": "x.pdf",
            "imported_at": datetime.now().isoformat(timespec="seconds"),
            "order_number": "REHYDR-2",
            "customer": "ACME",
            "status": "success",
            "portal_status": "found_in_fire",
            "snapshot": {
                "header": {"order_number": "REHYDR-2", "customer_name": "ACME"},
                "items": [{"description": "item A", "quantity": 3.0}],
                "source_file": "",
            },
            "check": {
                "available": False,
                "reason": "FB_DATABASE_NOT_SET",
                "client": {"match": False},
                "items": [],
                "summary": {},
            },
        }
    )
    repo.update_fire_poll_result(
        entry_id, "CANCELADO", datetime.now().isoformat(timespec="seconds")
    )

    r = client.get(f"/api/imported/{entry_id}/preview")
    assert r.status_code == 200
    body = r.json()
    assert body["portal_status"] == "found_in_fire"
    assert body["fire_status_last_seen"] == "CANCELADO"


def test_rehydrate_preview_fire_status_last_seen_absent_when_never_polled():
    """Sem poll ainda, o campo vem `None` — não deve virar string vazia nem
    quebrar o payload."""
    import uuid
    from datetime import datetime

    from app.persistence import repo

    entry_id = str(uuid.uuid4())
    repo.insert_import(
        {
            "id": entry_id,
            "source_filename": "x.pdf",
            "imported_at": datetime.now().isoformat(timespec="seconds"),
            "order_number": "REHYDR-3",
            "customer": "ACME",
            "status": "success",
            "portal_status": "parsed",
            "snapshot": {
                "header": {"order_number": "REHYDR-3", "customer_name": "ACME"},
                "items": [{"description": "item A", "quantity": 3.0}],
                "source_file": "",
            },
            "check": {
                "available": False,
                "reason": "FB_DATABASE_NOT_SET",
                "client": {"match": False},
                "items": [],
                "summary": {},
            },
        }
    )

    r = client.get(f"/api/imported/{entry_id}/preview")
    assert r.status_code == 200
    assert r.json()["fire_status_last_seen"] is None


def test_imported_filter_by_search_and_status():
    # Seed a few rows directly
    import uuid
    from datetime import datetime

    from app.persistence import repo

    for customer, status in [
        ("Riachuelo SA", "success"),
        ("Beira Rio", "success"),
        ("Erro LTDA", "error"),
    ]:
        repo.insert_import(
            {
                "id": str(uuid.uuid4()),
                "source_filename": "x.pdf",
                "imported_at": datetime.now().isoformat(timespec="seconds"),
                "order_number": "X-1",
                "customer": customer,
                "status": status,
                "output_files": [],
                "error": None if status == "success" else "boom",
            }
        )

    r = client.get("/api/imported?q=Riachuelo")
    assert r.status_code == 200
    body = r.json()
    assert body["total"] == 1
    assert body["entries"][0]["customer"] == "Riachuelo SA"

    r2 = client.get("/api/imported?status=error")
    assert r2.status_code == 200
    assert r2.json()["total"] == 1


def test_imported_filter_by_single_portal_status():
    """Um único `portal_status` na query continua funcionando exatamente como
    antes de `portal_status` virar `list[str] | None = Query(None)` — vira
    uma lista de 1 item, que `repo._build_where` trata como `IN (?)`."""
    import uuid
    from datetime import datetime

    from app.persistence import repo

    for order_number, status in [
        ("PS-1", "parsed"),
        ("PS-2", "sent_to_fire"),
        ("PS-3", "cancelled"),
    ]:
        repo.insert_import(
            {
                "id": str(uuid.uuid4()),
                "source_filename": "x.pdf",
                "imported_at": datetime.now().isoformat(timespec="seconds"),
                "order_number": order_number,
                "customer": "ACME",
                "status": "success",
                "portal_status": status,
            }
        )

    r = client.get("/api/imported?portal_status=parsed")
    assert r.status_code == 200
    body = r.json()
    assert body["total"] == 1
    assert body["entries"][0]["order_number"] == "PS-1"
    assert body["entries"][0]["portal_status"] == "parsed"


def test_imported_filter_by_multiple_portal_status():
    """`?portal_status=a&portal_status=b` (parâmetro repetido) filtra pela
    UNIÃO dos dois estados numa chamada só — é o que o chip "No Fire" da UI
    usa para casar `sent_to_fire` e `found_in_fire` juntos."""
    import uuid
    from datetime import datetime

    from app.persistence import repo

    for order_number, status in [
        ("PM-1", "sent_to_fire"),
        ("PM-2", "found_in_fire"),
        ("PM-3", "parsed"),
        ("PM-4", "cancelled"),
    ]:
        repo.insert_import(
            {
                "id": str(uuid.uuid4()),
                "source_filename": "x.pdf",
                "imported_at": datetime.now().isoformat(timespec="seconds"),
                "order_number": order_number,
                "customer": "ACME",
                "status": "success",
                "portal_status": status,
            }
        )

    r = client.get("/api/imported?portal_status=sent_to_fire&portal_status=found_in_fire")
    assert r.status_code == 200
    body = r.json()
    assert body["total"] == 2
    order_numbers = {e["order_number"] for e in body["entries"]}
    assert order_numbers == {"PM-1", "PM-2"}


def test_imported_counts_conta_todo_status_e_respeita_os_outros_filtros():
    """`counts` alimenta os contadores dos chips de status.

    Contrato em uma frase: conta por `portal_status` com TODOS os outros
    filtros aplicados e o próprio `portal_status` de fora. É o que faz o
    número do chip ser exatamente o que aparece ao clicar nele.

    Sem isso a reconciliação move a maioria dos pedidos de uma vez e a tela
    só encolhe — o operador vê "12 pedido(s)" e nenhuma pista de onde foram
    parar os outros 296.
    """
    import uuid
    from datetime import datetime

    from app.persistence import repo

    marca = f"CNT{uuid.uuid4().hex[:8].upper()}"
    for sufixo, status in [
        ("A", "parsed"),
        ("B", "sent_to_fire"),
        ("C", "found_in_fire"),
        ("D", "cancelled"),
        ("E", "error"),
    ]:
        repo.insert_import(
            {
                "id": str(uuid.uuid4()),
                "source_filename": "x.pdf",
                "imported_at": datetime.now().isoformat(timespec="seconds"),
                "order_number": f"{marca}-{sufixo}",
                "status": "success",
                "portal_status": status,
            }
        )

    r = client.get(f"/api/imported?q={marca}&portal_status=parsed")
    assert r.status_code == 200
    body = r.json()

    # `total` continua obedecendo o chip ativo — é o tamanho da lista.
    assert body["total"] == 1

    # `counts` NÃO obedece o chip ativo (é o eixo contado), mas obedece a
    # busca: só os cinco pedidos desta marca entram.
    assert body["counts"] == {
        "parsed": 1,
        "sent_to_fire": 1,
        "found_in_fire": 1,
        "cancelled": 1,
        "error": 1,
    }, "counts deve ignorar o filtro de portal_status e respeitar q"

    # O chip "Tudo" soma o dict; `error` não tem chip próprio mas aparece ao
    # clicar em "Tudo", então precisa estar na soma.
    assert sum(body["counts"].values()) == 5


# ── Manual cliente override (CLIENT_NOT_FOUND recovery) ──────────────────────


def _seed_parsed_entry(**overrides):
    import uuid
    from datetime import datetime

    from app.persistence import repo

    entry_id = overrides.pop("id", str(uuid.uuid4()))
    repo.insert_import(
        {
            "id": entry_id,
            "source_filename": "ovr.pdf",
            "imported_at": datetime.now().isoformat(timespec="seconds"),
            "order_number": overrides.get("order_number", "OVR-1"),
            "customer": overrides.get("customer", "ACME"),
            "customer_cnpj": overrides.get("customer_cnpj", "11.222.333/0001-44"),
            "status": "success",
            "portal_status": overrides.get("portal_status", "parsed"),
            "snapshot": {
                "header": {
                    "order_number": "OVR-1",
                    "customer_name": "ACME",
                    "customer_cnpj": "11.222.333/0001-44",
                },
                "items": [{"description": "x", "quantity": 1.0}],
                "source_file": "",
            },
            "check": {
                "available": True,
                "reason": None,
                "client": {
                    "match": False,
                    "fire_id": None,
                    "razao_social": None,
                    "cnpj": "11.222.333/0001-44",
                },
                "items": [],
                "summary": {
                    "items_total": 1,
                    "items_matched": 0,
                    "items_missing": 1,
                    "client_matched": False,
                },
            },
        }
    )
    return entry_id


class _FakeFbCursor:
    """Minimal cursor double — replays a list of fetchone() rows in order."""

    def __init__(self, rows):
        self._rows = list(rows)
        self.executed: list[tuple] = []

    def execute(self, sql, params=()):
        self.executed.append((sql, params))

    def fetchone(self):
        return self._rows.pop(0) if self._rows else None

    def fetchall(self):
        out, self._rows = self._rows, []
        return out

    def close(self):
        pass


def _patch_fb_with_rows(monkeypatch, rows, *, configured=True):
    """Patch FirebirdConnection so .is_configured() and .connect() return our fake."""
    from app.erp import connection as conn_mod

    def _ctor():
        inst = conn_mod.FirebirdConnection.__new__(conn_mod.FirebirdConnection)
        return inst

    cursor = _FakeFbCursor(rows)

    class _CtxConn:
        def cursor(self):
            return cursor

        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

    inst = _ctor()
    monkeypatch.setattr(inst, "is_configured", lambda: configured, raising=False)
    monkeypatch.setattr(inst, "connect", lambda: _CtxConn(), raising=False)

    # Patch the constructor so server.py's `FirebirdConnection()` returns our instance
    monkeypatch.setattr(conn_mod, "FirebirdConnection", lambda: inst)
    return cursor


def test_search_clientes_requires_min_length():
    r = client.get("/api/clientes/search?q=a")
    assert r.status_code == 400


def test_search_clientes_returns_503_when_fb_not_configured(monkeypatch):
    _patch_fb_with_rows(monkeypatch, [], configured=False)
    r = client.get("/api/clientes/search?q=acme")
    assert r.status_code == 503


def test_search_clientes_happy_path_returns_results(monkeypatch):
    cursor = _patch_fb_with_rows(
        monkeypatch,
        [
            (101, "ACME COMERCIO LTDA", "11.222.333/0001-44"),
            (102, "ACME INDUSTRIAS SA", "55.666.777/0001-88"),
        ],
    )
    r = client.get("/api/clientes/search?q=acme")
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["total_returned"] == 2
    assert body["results"][0]["codigo"] == 101
    assert body["results"][0]["razao_social"] == "ACME COMERCIO LTDA"
    # Must have run SEARCH_CLIENTS with razao + cnpj patterns.
    sql, params = cursor.executed[0]
    assert "CADASTRO" in sql and "RAZAO_SOCIAL" in sql
    assert params[0] == "%ACME%"


def test_search_clientes_uses_selected_environment_firebird(monkeypatch, tmp_path):
    from app.erp import connection as conn_mod
    from app.persistence import environments_repo

    env = environments_repo.create(
        slug="mm",
        name="MM",
        watch_dir=str(tmp_path / "in"),
        output_dir=str(tmp_path / "out"),
        fb_path=str(tmp_path / "mm.fdb"),
        fb_host="127.0.0.1",
        fb_port="3051",
        fb_user="SYSDBA",
        fb_charset="WIN1252",
        fb_password="secret",
    )
    cursor = _FakeFbCursor([(101, "ACME COMERCIO LTDA", "11.222.333/0001-44")])
    seen_cfg: dict = {}

    class _CtxConn:
        def cursor(self):
            return cursor

        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

    inst = conn_mod.FirebirdConnection.__new__(conn_mod.FirebirdConnection)
    monkeypatch.setattr(inst, "is_configured", lambda: False, raising=False)
    monkeypatch.setattr(
        inst,
        "connect",
        lambda: (_ for _ in ()).throw(AssertionError("legacy connect must not be used")),
        raising=False,
    )

    def _connect_with_config(cfg):
        seen_cfg.update(cfg)
        return _CtxConn()

    monkeypatch.setattr(inst, "connect_with_config", _connect_with_config, raising=False)
    monkeypatch.setattr(conn_mod, "FirebirdConnection", lambda: inst)

    client.cookies.set("portal_env", env["id"])
    try:
        r = client.get("/api/clientes/search?q=acme")
    finally:
        client.cookies.clear()
    assert r.status_code == 200, r.text
    assert r.json()["results"][0]["codigo"] == 101
    assert seen_cfg["path"] == str(tmp_path / "mm.fdb")
    assert seen_cfg["host"] == "127.0.0.1"


def test_search_clientes_strips_non_digits_for_cnpj_pattern(monkeypatch):
    cursor = _patch_fb_with_rows(monkeypatch, [])
    r = client.get("/api/clientes/search?q=11.222.333/0001-44")
    assert r.status_code == 200
    _sql, params = cursor.executed[0]
    assert params[1] == "%11222333000144%"


def test_override_cliente_happy_path(monkeypatch):
    from app.persistence import repo

    entry_id = _seed_parsed_entry()
    _patch_fb_with_rows(
        monkeypatch,
        [
            (4242, "ACME COMERCIO LTDA", "11222333000144"),  # FIND_CLIENT_BY_CODIGO
        ],
    )

    r = client.post(
        f"/api/imported/{entry_id}/override-cliente",
        json={"cliente_codigo": 4242, "reason": "varejista mudou de razão social"},
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["cliente_override_codigo"] == 4242
    assert body["cliente_override_razao"] == "ACME COMERCIO LTDA"

    got = repo.get_import(entry_id)
    assert got["cliente_override_codigo"] == 4242
    assert got["cliente_override_razao"] == "ACME COMERCIO LTDA"
    assert got["cliente_override_at"]
    # TEST_AUTH_BYPASS substitui require_user pelo _TEST_USER (test@portal.local).
    assert got["cliente_override_by"] == "test@portal.local"


def test_override_cliente_appends_audit_with_actor_email(monkeypatch):
    from app.persistence import repo

    entry_id = _seed_parsed_entry()
    _patch_fb_with_rows(monkeypatch, [(4242, "ACME LTDA", "11222333000144")])

    client.post(
        f"/api/imported/{entry_id}/override-cliente",
        json={"cliente_codigo": 4242, "reason": "manual fix"},
    )
    events = repo.list_audit(entry_id)
    e = next(ev for ev in events if ev["event_type"] == "cliente_override_selected")
    assert e["detail"]["cliente_codigo"] == 4242
    assert e["detail"]["cliente_razao"] == "ACME LTDA"
    assert e["detail"]["reason"] == "manual fix"
    assert e["detail"]["user_email"] == "test@portal.local"
    assert e["detail"]["user_id"] == 0
    assert e["detail"]["previous_cnpj"] == "11.222.333/0001-44"


def test_override_cliente_rejects_non_parsed(monkeypatch):
    entry_id = _seed_parsed_entry(portal_status="cancelled")
    r = client.post(
        f"/api/imported/{entry_id}/override-cliente",
        json={"cliente_codigo": 4242},
    )
    assert r.status_code == 409
    assert "revisão" in r.json()["detail"].lower()


def test_override_cliente_rejects_unknown_codigo(monkeypatch):
    entry_id = _seed_parsed_entry()
    _patch_fb_with_rows(monkeypatch, [None])  # FIND_CLIENT_BY_CODIGO sem match
    r = client.post(
        f"/api/imported/{entry_id}/override-cliente",
        json={"cliente_codigo": 99999999},
    )
    assert r.status_code == 422


def test_override_cliente_returns_404_when_entry_missing():
    r = client.post(
        "/api/imported/does-not-exist/override-cliente",
        json={"cliente_codigo": 1},
    )
    assert r.status_code == 404


def test_override_cliente_returns_503_when_fb_not_configured(monkeypatch):
    entry_id = _seed_parsed_entry()
    _patch_fb_with_rows(monkeypatch, [], configured=False)
    r = client.post(
        f"/api/imported/{entry_id}/override-cliente",
        json={"cliente_codigo": 4242},
    )
    assert r.status_code == 503


def test_rehydrate_preview_injects_override_into_check_block(monkeypatch):
    from app.persistence import repo

    entry_id = _seed_parsed_entry()
    repo.set_client_override(entry_id, codigo=4242, razao="ACME OVERRIDE LTDA")

    r = client.get(f"/api/imported/{entry_id}/preview")
    assert r.status_code == 200
    body = r.json()
    assert body["check"]["client"]["match"] is True
    assert body["check"]["client"]["fire_id"] == 4242
    assert body["check"]["client"]["override"] is True
    assert body["check"]["summary"]["client_matched"] is True
    assert body["cliente_override"]["codigo"] == 4242
    assert body["cliente_override"]["razao_social"] == "ACME OVERRIDE LTDA"
    assert body["cliente_override"]["by"] is None


def test_rehydrate_preview_omits_override_block_when_no_override():
    entry_id = _seed_parsed_entry()
    r = client.get(f"/api/imported/{entry_id}/preview")
    assert r.status_code == 200
    body = r.json()
    assert body["cliente_override"] is None
    # Original false-match preserved
    assert body["check"]["client"]["match"] is False


def test_send_to_fire_passes_override_to_exporter(monkeypatch):
    """Override stored em imports é lido e passado ao FirebirdExporter."""
    from app import config as app_config
    from app.exporters import firebird_exporter as fb_mod
    from app.persistence import repo

    entry_id = _seed_parsed_entry()
    repo.set_client_override(entry_id, codigo=4242, razao="ACME OVERRIDE")

    captured = {}

    def _fake_export(self, order, *, override_client_id=None, produtos_por_item=None):
        captured["override"] = override_client_id
        return fb_mod.FirebirdExportResult(
            order_number=order.header.order_number,
            items_inserted=1,
            fire_codigo=999,
        )

    monkeypatch.setattr(fb_mod.FirebirdExporter, "export", _fake_export)
    monkeypatch.setattr(
        app_config,
        "load",
        lambda: {
            "watch_dir": ".",
            "output_dir": ".",
            "export_mode": "db",
        },
    )

    r = client.post(f"/api/imported/{entry_id}/send-to-fire")
    assert r.status_code == 200, r.text
    assert captured["override"] == 4242
    assert r.json()["fire_codigo"] == 999


@pytest.mark.skipif(not SAMPLES.exists(), reason="samples/ directory not found")
def test_process_sbf_centauro(tmp_path):
    pdf = SAMPLES / "Pedido_0029852483.pdf"
    if not pdf.exists():
        pytest.skip("Sample PDF not available")

    r = client.post(
        "/api/process",
        data={"output_dir": str(tmp_path)},
        files=[("files", (pdf.name, pdf.read_bytes(), "application/pdf"))],
    )
    assert r.status_code == 200
    data = r.json()
    assert len(data["results"]) == 1
    assert data["results"][0]["order"] == "29852483"


# ── Regressão: salvar config preserva sessão ─────────────────────────────────
# Antes do fix, o `app_state.db` era resolvido a partir de `watch_dir`, então
# `POST /api/config` movia o arquivo do banco e deixava a sessão órfã (cookie
# válido apontando para token em DB que o app não enxergava mais). O fix em
# `app/persistence/db.py` pinou o caminho em local estável independente de
# `watch_dir`.


def test_save_config_preserves_session(real_auth, tmp_path, monkeypatch):
    from fastapi.testclient import TestClient

    from app import config as app_config
    from app.persistence import users_repo
    from app.web.server import app as fastapi_app

    monkeypatch.setenv("RATE_LIMIT_ENABLED", "false")
    # Isolar config.json para não contaminar outros testes.
    monkeypatch.setattr(app_config, "_CONFIG_FILE", tmp_path / "config.json")
    db.set_db_path(tmp_path / "app_state.db")
    db.reset_init_cache()
    db.init()

    users_repo.create_user(
        email="cfg@portal.local",
        password="strongpass1",
        role="admin",
    )

    c = TestClient(fastapi_app)
    r = c.post(
        "/api/auth/login",
        json={"email": "cfg@portal.local", "password": "strongpass1"},
    )
    assert r.status_code == 200, r.text
    cookie = r.cookies.get("portal_session")
    assert cookie

    headers = {"Cookie": f"portal_session={cookie}"}

    for i in range(3):
        new_watch = tmp_path / f"watch{i}"
        new_out = tmp_path / f"out{i}"
        r = c.post(
            "/api/config",
            json={
                "watchDir": str(new_watch),
                "outputDir": str(new_out),
                "exportMode": "xlsx",
            },
            headers=headers,
        )
        assert r.status_code == 200, f"save #{i} falhou: {r.text}"

    me = c.get("/api/auth/me", headers=headers)
    assert me.status_code == 200, me.text
    assert me.json()["user"]["email"] == "cfg@portal.local"


# ── ack-sem-preco ─────────────────────────────────────────────────────────────


def test_ack_sem_preco_persists_and_audits(monkeypatch):
    """POST /api/imported/{id}/ack-sem-preco grava sidecar + audit."""
    import uuid
    from datetime import datetime

    from app.erp import product_check as pc_mod
    from app.persistence import repo

    entry_id = str(uuid.uuid4())
    repo.insert_import(
        {
            "id": entry_id,
            "source_filename": "x.pdf",
            "imported_at": datetime.now().isoformat(timespec="seconds"),
            "order_number": "ACK-1",
            "status": "success",
            "portal_status": "parsed",
            "snapshot": {
                "header": {"order_number": "ACK-1", "customer_cnpj": "00000000000100"},
                "items": [
                    {"description": "x", "quantity": 1.0, "ean": "7891", "unit_price": 89.90}
                ],
                "source_file": "",
            },
        }
    )

    fake_check = {
        "available": True,
        "items": [
            {
                "ean": "7891",
                "product_code": None,
                "price_status": "no_price_in_fire",
                "fire_product_id": 42,
            }
        ],
        "summary": {},
    }
    monkeypatch.setattr(pc_mod, "check_order", lambda order, **kw: fake_check)

    r = client.post(f"/api/imported/{entry_id}/ack-sem-preco")
    assert r.status_code == 200, r.text

    body = r.json()
    assert body["ack_by"] == "test@portal.local"  # do TEST_AUTH_BYPASS
    assert body["items_acked"] == [
        {"ean": "7891", "product_code": None, "fire_product_id": 42},
    ]

    got = repo.get_import(entry_id)
    assert got["sem_preco_ack_by"] == "test@portal.local"
    assert got["sem_preco_ack_items"] == [
        {"ean": "7891", "product_code": None, "fire_product_id": 42},
    ]

    audits = [a["event_type"] for a in repo.list_audit(entry_id)]
    assert "sem_preco_acknowledged" in audits


def test_ack_sem_preco_rejects_wrong_status():
    import uuid
    from datetime import datetime

    from app.persistence import repo

    entry_id = str(uuid.uuid4())
    repo.insert_import(
        {
            "id": entry_id,
            "source_filename": "x.pdf",
            "imported_at": datetime.now().isoformat(timespec="seconds"),
            "order_number": "ACK-2",
            "status": "success",
            "portal_status": "sent_to_fire",
            "snapshot": {"header": {"order_number": "ACK-2"}, "items": []},
        }
    )
    r = client.post(f"/api/imported/{entry_id}/ack-sem-preco")
    assert r.status_code == 409


def test_ack_sem_preco_returns_404_when_missing():
    r = client.post("/api/imported/does-not-exist/ack-sem-preco")
    assert r.status_code == 404


def test_rehydrate_preview_surfaces_sem_preco_ack():
    import uuid
    from datetime import datetime

    from app.persistence import repo

    entry_id = str(uuid.uuid4())
    repo.insert_import(
        {
            "id": entry_id,
            "source_filename": "x.pdf",
            "imported_at": datetime.now().isoformat(timespec="seconds"),
            "order_number": "REHY-ACK",
            "status": "success",
            "portal_status": "parsed",
            "snapshot": {
                "header": {"order_number": "REHY-ACK"},
                "items": [{"description": "x", "quantity": 1.0}],
                "source_file": "",
            },
        }
    )
    repo.set_sem_preco_ack(
        entry_id,
        by_email="op@example.com",
        items=[{"ean": "7891", "product_code": None, "fire_product_id": 1}],
    )

    r = client.get(f"/api/imported/{entry_id}/preview")
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["sem_preco_ack"]["by"] == "op@example.com"
    assert body["sem_preco_ack"]["items"][0]["ean"] == "7891"


# ── Guard _send_one_to_fire: price-check defense in depth ────────────────────


def _seed_parsed_order(entry_id, *, items=None, snapshot_items=None):
    """Helper: cria entry parsed pronto para send-to-fire."""
    from datetime import datetime

    from app.persistence import repo

    repo.insert_import(
        {
            "id": entry_id,
            "source_filename": "p.pdf",
            "imported_at": datetime.now().isoformat(timespec="seconds"),
            "order_number": f"GUARD-{entry_id[:4]}",
            "customer": "ACME",
            "status": "success",
            "portal_status": "parsed",
            "snapshot": {
                "header": {"order_number": f"GUARD-{entry_id[:4]}", "customer_name": "ACME"},
                "items": snapshot_items
                or [{"description": "x", "quantity": 1.0, "ean": "7891", "unit_price": 89.90}],
                "source_file": "",
            },
        }
    )


def test_send_to_fire_blocked_by_price_mismatch(monkeypatch):
    import uuid

    from app.erp import product_check as pc_mod
    from app.exporters import firebird_exporter as fb_mod
    from app.persistence import repo

    entry_id = str(uuid.uuid4())
    _seed_parsed_order(entry_id)

    fake_check = {
        "available": True,
        "items": [
            {
                "ean": "7891",
                "product_code": None,
                "price_status": "mismatch",
                "unit_price_order": 89.90,
                "fire_preco_venda": 90.00,
            }
        ],
        "summary": {},
    }
    monkeypatch.setattr(pc_mod, "check_order", lambda order, **kw: fake_check)

    called = {"export": False}

    def _no_export(self, order, *, override_client_id=None, produtos_por_item=None):
        called["export"] = True
        raise AssertionError("FirebirdExporter.export não pode ser chamado")

    monkeypatch.setattr(fb_mod.FirebirdExporter, "export", _no_export)

    r = client.post(f"/api/imported/{entry_id}/send-to-fire")
    assert r.status_code == 409
    assert "preço" in r.json()["detail"].lower() or "price" in r.json()["detail"].lower()
    assert called["export"] is False

    audits = [a["event_type"] for a in repo.list_audit(entry_id)]
    assert "send_to_fire_blocked" in audits


def test_send_to_fire_blocked_by_no_price_unacked(monkeypatch):
    import uuid

    from app.erp import product_check as pc_mod
    from app.exporters import firebird_exporter as fb_mod

    entry_id = str(uuid.uuid4())
    _seed_parsed_order(entry_id)

    fake_check = {
        "available": True,
        "items": [{"ean": "7891", "product_code": None, "price_status": "no_price_in_fire"}],
        "summary": {},
    }
    monkeypatch.setattr(pc_mod, "check_order", lambda order, **kw: fake_check)
    monkeypatch.setattr(
        fb_mod.FirebirdExporter,
        "export",
        lambda *a, **kw: (_ for _ in ()).throw(AssertionError("não pode chamar")),
    )

    r = client.post(f"/api/imported/{entry_id}/send-to-fire")
    assert r.status_code == 409


def test_send_to_fire_passes_with_ack(monkeypatch):
    import uuid

    from app import config as app_config
    from app.erp import product_check as pc_mod
    from app.exporters import firebird_exporter as fb_mod
    from app.persistence import repo

    entry_id = str(uuid.uuid4())
    _seed_parsed_order(entry_id)
    repo.set_sem_preco_ack(
        entry_id, by_email="op@example.com", items=[{"ean": "7891", "product_code": None}]
    )

    fake_check = {
        "available": True,
        "items": [{"ean": "7891", "product_code": None, "price_status": "no_price_in_fire"}],
        "summary": {},
    }
    monkeypatch.setattr(pc_mod, "check_order", lambda order, **kw: fake_check)
    monkeypatch.setattr(
        fb_mod.FirebirdExporter,
        "export",
        lambda self, order, **kw: fb_mod.FirebirdExportResult(
            order_number=order.header.order_number,
            items_inserted=1,
            fire_codigo=999,
        ),
    )
    monkeypatch.setattr(
        app_config, "load", lambda: {"watch_dir": ".", "output_dir": ".", "export_mode": "db"}
    )

    r = client.post(f"/api/imported/{entry_id}/send-to-fire")
    assert r.status_code == 200, r.text
    assert r.json()["fire_codigo"] == 999


def test_send_to_fire_passes_when_check_unavailable(monkeypatch):
    """Fire offline → check best-effort, não bloqueia."""
    import uuid

    from app import config as app_config
    from app.erp import product_check as pc_mod
    from app.exporters import firebird_exporter as fb_mod

    entry_id = str(uuid.uuid4())
    _seed_parsed_order(entry_id)

    monkeypatch.setattr(
        pc_mod, "check_order", lambda order, **kw: {"available": False, "items": [], "summary": {}}
    )
    monkeypatch.setattr(
        fb_mod.FirebirdExporter,
        "export",
        lambda self, order, **kw: fb_mod.FirebirdExportResult(
            order_number=order.header.order_number,
            items_inserted=1,
            fire_codigo=111,
        ),
    )
    monkeypatch.setattr(
        app_config, "load", lambda: {"watch_dir": ".", "output_dir": ".", "export_mode": "db"}
    )

    r = client.post(f"/api/imported/{entry_id}/send-to-fire")
    assert r.status_code == 200
    assert r.json()["fire_codigo"] == 111


# ── Guard _export_one_xlsx: price-check defense in depth ─────────────────────


def test_export_xlsx_blocked_by_price_mismatch(monkeypatch):
    import uuid

    from app.erp import product_check as pc_mod
    from app.persistence import repo

    entry_id = str(uuid.uuid4())
    _seed_parsed_order(entry_id)

    fake_check = {
        "available": True,
        "items": [
            {
                "ean": "7891",
                "product_code": None,
                "price_status": "mismatch",
                "unit_price_order": 89.90,
                "fire_preco_venda": 90.00,
            }
        ],
        "summary": {},
    }
    monkeypatch.setattr(pc_mod, "check_order", lambda order, **kw: fake_check)

    r = client.post(f"/api/imported/{entry_id}/export-xlsx")
    assert r.status_code == 409

    audits = [a["event_type"] for a in repo.list_audit(entry_id)]
    assert "xlsx_export_blocked" in audits


def test_export_xlsx_passes_with_ack(monkeypatch, tmp_path):
    import uuid

    from app import config as app_config
    from app.erp import product_check as pc_mod
    from app.persistence import repo

    entry_id = str(uuid.uuid4())
    _seed_parsed_order(entry_id)
    repo.set_sem_preco_ack(
        entry_id, by_email="op@example.com", items=[{"ean": "7891", "product_code": None}]
    )

    fake_check = {
        "available": True,
        "items": [{"ean": "7891", "product_code": None, "price_status": "no_price_in_fire"}],
        "summary": {},
    }
    monkeypatch.setattr(pc_mod, "check_order", lambda order, **kw: fake_check)
    monkeypatch.setattr(
        app_config,
        "load",
        lambda: {"watch_dir": str(tmp_path), "output_dir": str(tmp_path), "export_mode": "xlsx"},
    )

    r = client.post(f"/api/imported/{entry_id}/export-xlsx")
    assert r.status_code == 200, r.text


def _insert_parsed_import(entry_id: str, order_number: str = "PUSH-1"):
    import uuid as _uuid  # noqa: F401
    from datetime import datetime

    from app.persistence import repo

    repo.insert_import(
        {
            "id": entry_id,
            "source_filename": "pedido.pdf",
            "imported_at": datetime.now().isoformat(timespec="seconds"),
            "order_number": order_number,
            "customer": "ACME",
            "status": "success",
            "portal_status": "parsed",
            "snapshot": {
                "header": {"order_number": order_number, "customer_name": "ACME"},
                "items": [{"description": "x", "quantity": 1.0}],
                "source_file": "",
            },
        }
    )


def test_export_xlsx_dispara_push_flowpcp(monkeypatch, tmp_path):
    """Gerar XLS (modo xlsx) TAMBÉM empurra o pedido pro FlowPCP — Fire fica
    manual (XLS) mas o Flow já recebe. Best-effort: gated por flowpcp_enabled
    dentro do hook; aqui provamos que o hook é chamado com o slug do ambiente."""
    import uuid

    from app.erp import product_check as pc_mod
    from app.web import server as server_mod

    entry_id = str(uuid.uuid4())
    _insert_parsed_import(entry_id)

    # Env fake não tem Firebird — stub do price-check (não é o que se testa aqui).
    monkeypatch.setattr(
        pc_mod,
        "check_order",
        lambda order, **kw: {"available": False, "items": [], "summary": {}},
    )

    pushes: list[tuple[str, str]] = []
    monkeypatch.setattr(
        server_mod,
        "push_new_order",
        lambda order, *, import_id, slug: pushes.append((import_id, slug)) or True,
    )

    cfg = {"watch_dir": str(tmp_path), "output_dir": str(tmp_path), "export_mode": "xlsx"}
    out = server_mod._export_one_xlsx(entry_id, cfg, request_env={"slug": "mm"})
    assert out.ok, out.detail
    assert pushes == [(entry_id, "mm")]


def test_export_xlsx_sem_ambiente_nao_chama_push(monkeypatch, tmp_path):
    import uuid

    from app.web import server as server_mod

    entry_id = str(uuid.uuid4())
    _insert_parsed_import(entry_id, order_number="PUSH-2")

    pushes = []
    monkeypatch.setattr(
        server_mod,
        "push_new_order",
        lambda order, *, import_id, slug: pushes.append(slug) or True,
    )
    cfg = {"watch_dir": str(tmp_path), "output_dir": str(tmp_path), "export_mode": "xlsx"}
    out = server_mod._export_one_xlsx(entry_id, cfg, request_env=None)
    assert out.ok
    assert pushes == []


# ── De-para de produto: busca local, vínculo, undo ────────────────────────


def _seed_catalogo_fire(rows: list[dict]) -> None:
    """rows: [{"fire_produto_id", "codigo", "nome", "ean", "unidade"?, "ativo"?, "tipo"?}]."""
    from app.erp.catalog_extract import ProdutoFireDTO
    from app.persistence import catalogo_fire_repo, db

    dtos = [
        ProdutoFireDTO(
            fire_produto_id=r["fire_produto_id"],
            codigo=r["codigo"],
            nome=r["nome"],
            unidade=r.get("unidade", "PAR"),
            ean=r.get("ean"),
            ativo=r.get("ativo", True),
            tipo=r.get("tipo", "simples"),
        )
        for r in rows
    ]
    with db.connect() as conn:
        catalogo_fire_repo.replace_all(conn, dtos, extraido_em="2026-07-25T00:00:00")


def _seed_parsed_import_with_item(entry_id: str, *, item: dict, customer_cnpj: str) -> None:
    from datetime import datetime

    from app.persistence import repo

    repo.insert_import(
        {
            "id": entry_id,
            "source_filename": "pedido.pdf",
            "imported_at": datetime.now().isoformat(timespec="seconds"),
            "order_number": "DP-1",
            "customer": "ACME",
            "customer_cnpj": customer_cnpj,
            "status": "success",
            "portal_status": "parsed",
            "snapshot": {
                "header": {
                    "order_number": "DP-1",
                    "customer_name": "ACME",
                    "customer_cnpj": customer_cnpj,
                },
                "items": [item],
                "source_file": "",
            },
        }
    )


def test_produtos_search_local():
    _seed_catalogo_fire(
        [{"fire_produto_id": "10", "codigo": "10", "nome": "TENIS AZUL", "ean": "789"}]
    )
    r = client.get("/api/produtos/search?q=azul")
    assert r.status_code == 200, r.text
    body = r.json()
    assert any(p["fire_codigo"] == "10" for p in body["results"])
    assert any(p["fire_nome"] == "TENIS AZUL" for p in body["results"])


def test_produtos_search_requires_min_length():
    r = client.get("/api/produtos/search?q=a")
    assert r.status_code == 400


def test_produtos_search_hints_only_no_q():
    """Picker de de-para abre sem q, só com dicas — deve popular suggestions."""
    _seed_catalogo_fire(
        [
            {"fire_produto_id": "10", "codigo": "10", "nome": "TENIS AZUL CORRIDA", "ean": "789"},
            {"fire_produto_id": "11", "codigo": "11", "nome": "SANDALIA PRETA", "ean": "456"},
        ]
    )
    r = client.get("/api/produtos/search?desc=tenis%20azul%20corrida&code=10")
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["results"] == []
    assert body["suggestions"], "esperava suggestions não vazio mesmo sem q"
    assert body["suggestions"][0]["fire_produto_id"] == "10"
    scores = [s["score"] for s in body["suggestions"]]
    assert scores == sorted(scores, reverse=True)


def test_produtos_search_suggestions_ranked():
    _seed_catalogo_fire(
        [
            {"fire_produto_id": "10", "codigo": "10", "nome": "TENIS AZUL CORRIDA", "ean": "789"},
            {"fire_produto_id": "11", "codigo": "11", "nome": "SANDALIA PRETA", "ean": "456"},
        ]
    )
    r = client.get("/api/produtos/search?q=azul&desc=tenis%20azul")
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["suggestions"], "esperava suggestions não vazio"
    assert body["suggestions"][0]["fire_produto_id"] == "10"
    assert body["suggestions"][0]["score"] > 0
    scores = [s["score"] for s in body["suggestions"]]
    assert scores == sorted(scores, reverse=True)


def test_vincular_produto_persiste():
    import uuid

    from app.persistence import catalogo_fire_repo, db, produto_depara_repo

    entry_id = str(uuid.uuid4())
    cnpj = "11.222.333/0001-44"
    _seed_parsed_import_with_item(
        entry_id,
        item={
            "description": "TENIS SEM MATCH",
            "quantity": 1.0,
            "product_code": "ABC123",
            "ean": "789",
        },
        customer_cnpj=cnpj,
    )
    _seed_catalogo_fire(
        [{"fire_produto_id": "10", "codigo": "10", "nome": "TENIS AZUL", "ean": "789"}]
    )

    r = client.post(
        f"/api/imported/{entry_id}/vincular-produto",
        json={"item_index": 0, "fire_produto_id": "10"},
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["entry_id"] == entry_id
    assert "check" in body
    chaves = {(v["chave_tipo"], v["chave_valor"]) for v in body["vinculos"]}
    assert ("codigo", "ABC123") in chaves
    assert ("ean", "789") in chaves

    ckey = produto_depara_repo.client_key(cnpj, "ACME")
    with db.connect() as conn:
        resolved = produto_depara_repo.lookup(conn, ckey, codigos=["ABC123"], eans=["789"])
    assert resolved[("codigo", "ABC123")]["fire_produto_id"] == "10"
    assert resolved[("codigo", "ABC123")]["fire_codigo"] == "10"
    assert resolved[("ean", "789")]["fire_produto_id"] == "10"

    # sanity: catalogo_fire copy used to resolve the target product is present.
    with db.connect() as conn:
        assert any(p["fire_produto_id"] == "10" for p in catalogo_fire_repo.list_all(conn))


def test_vincular_rejects_wrong_status():
    entry_id = _seed_parsed_entry(portal_status="cancelled")
    r = client.post(
        f"/api/imported/{entry_id}/vincular-produto",
        json={"item_index": 0, "fire_produto_id": "10"},
    )
    assert r.status_code == 409


def test_vincular_produto_404_quando_pedido_nao_existe():
    r = client.post(
        "/api/imported/nao-existe/vincular-produto",
        json={"item_index": 0, "fire_produto_id": "10"},
    )
    assert r.status_code == 404


def test_vincular_produto_422_item_index_fora_do_intervalo():
    import uuid

    entry_id = str(uuid.uuid4())
    _seed_parsed_import_with_item(
        entry_id,
        item={
            "description": "TENIS SEM MATCH",
            "quantity": 1.0,
            "product_code": "ABC123",
            "ean": "789",
        },
        customer_cnpj="11.222.333/0001-44",
    )

    r = client.post(
        f"/api/imported/{entry_id}/vincular-produto",
        json={"item_index": 5, "fire_produto_id": "10"},
    )
    assert r.status_code == 422


def test_vincular_produto_422_fire_produto_id_desconhecido():
    import uuid

    entry_id = str(uuid.uuid4())
    _seed_parsed_import_with_item(
        entry_id,
        item={
            "description": "TENIS SEM MATCH",
            "quantity": 1.0,
            "product_code": "ABC123",
            "ean": "789",
        },
        customer_cnpj="11.222.333/0001-44",
    )
    _seed_catalogo_fire(
        [{"fire_produto_id": "10", "codigo": "10", "nome": "TENIS AZUL", "ean": "789"}]
    )

    r = client.post(
        f"/api/imported/{entry_id}/vincular-produto",
        json={"item_index": 0, "fire_produto_id": "999"},
    )
    assert r.status_code == 422


def test_vincular_produto_regrava_o_check_preco(monkeypatch):
    import uuid

    from app.erp import preco_trava
    from app.persistence import repo

    entry_id = str(uuid.uuid4())
    _seed_parsed_import_with_item(
        entry_id,
        item={"description": "TENIS", "quantity": 1.0, "product_code": "ABC123", "ean": "789"},
        customer_cnpj="11.222.333/0001-44",
    )
    _seed_catalogo_fire(
        [{"fire_produto_id": "10", "codigo": "10", "nome": "TENIS AZUL", "ean": "789"}]
    )
    monkeypatch.setattr(
        preco_trava, "check_preco_para_gravar", lambda order, env: {"available": True, "items": []}
    )

    r = client.post(
        f"/api/imported/{entry_id}/vincular-produto",
        json={"item_index": 0, "fire_produto_id": "10"},
    )

    assert r.status_code == 200, r.text
    assert repo.get_import(entry_id)["check_preco"] == {"available": True, "items": []}


def test_vincular_produto_sem_fonte_nao_grava_check_preco():
    import uuid

    from app.persistence import repo

    entry_id = str(uuid.uuid4())
    _seed_parsed_import_with_item(
        entry_id,
        item={"description": "TENIS", "quantity": 1.0, "product_code": "ABC123", "ean": "789"},
        customer_cnpj="11.222.333/0001-44",
    )
    _seed_catalogo_fire(
        [{"fire_produto_id": "10", "codigo": "10", "nome": "TENIS AZUL", "ean": "789"}]
    )

    r = client.post(
        f"/api/imported/{entry_id}/vincular-produto",
        json={"item_index": 0, "fire_produto_id": "10"},
    )

    assert r.status_code == 200, r.text
    assert repo.get_import(entry_id)["check_preco"] is None


def test_delete_depara_desfaz():
    from datetime import datetime

    from app.persistence import db, produto_depara_repo

    cnpj = "11.222.333/0001-44"
    ckey = produto_depara_repo.client_key(cnpj, None)
    with db.connect() as conn:
        produto_depara_repo.upsert(
            conn,
            client_key=ckey,
            chave_tipo="codigo",
            chave_valor="ABC123",
            fire_produto_id="10",
            fire_codigo="10",
            fire_ean="789",
            fire_nome="TENIS AZUL",
            criado_em=datetime.now().isoformat(timespec="seconds"),
            criado_por="test@portal.local",
        )
        dep_id = produto_depara_repo.list_for_client(conn, ckey)[0]["id"]

    r = client.delete(f"/api/produtos/depara/{dep_id}")
    assert r.status_code == 200, r.text
    assert r.json()["deleted"] == dep_id

    with db.connect() as conn:
        assert produto_depara_repo.list_for_client(conn, ckey) == []


# ── Arquivo original: cópia imutável de tudo que entra ───────────────────────
# Caso AF127/AF017 (H2S4, 27/07/2026): dois arquivos com o mesmo nome no mesmo
# dia, o segundo já errado, e nenhuma cópia nossa pra provar. Agora todo
# recebimento (upload, pasta vigiada, lote) guarda antes do parse.


def _recebidos_root() -> Path:
    import os

    return Path(os.environ["APP_DATA_DIR"]) / "recebidos"


def _arquivos_recebidos() -> list[Path]:
    root = _recebidos_root()
    if not root.exists():
        return []
    return sorted(p for p in root.rglob("*") if p.is_file())


def _sample_pdf() -> Path:
    pdf = SAMPLES / "2600009562-2026-02-25.pdf"
    if not pdf.exists():
        pytest.skip("Sample PDF not available")
    return pdf


def _config_em_tmp(monkeypatch, tmp_path: Path, watch: Path | None = None) -> Path:
    from app import config as app_config

    imp = tmp_path / "imported"
    imp.mkdir(exist_ok=True)
    monkeypatch.setattr(
        app_config,
        "load",
        lambda: {
            "watch_dir": str(watch or tmp_path),
            "output_dir": str(tmp_path),
            "export_mode": "xlsx",
        },
    )
    monkeypatch.setattr(app_config, "imported_dir", lambda _cfg: imp)
    return imp


def _insert_import_minimo(**extra) -> str:
    import uuid
    from datetime import datetime

    from app.persistence import repo

    entry = {
        "id": str(uuid.uuid4()),
        "source_filename": extra.pop("source_filename", "velho.pdf"),
        "imported_at": datetime.now().isoformat(timespec="seconds"),
        "order_number": "X1",
        "customer": "Cliente",
        "output_files": [],
        "status": "success",
        "snapshot": None,
        **extra,
    }
    repo.insert_import(entry)
    return entry["id"]


def test_preview_por_upload_guarda_copia_do_original_antes_do_parse():
    pdf = _sample_pdf()

    r = client.post(
        "/api/preview", files=[("file", (pdf.name, pdf.read_bytes(), "application/pdf"))]
    )

    assert r.status_code == 200, r.text
    copias = _arquivos_recebidos()
    assert len(copias) == 1
    assert copias[0].read_bytes() == pdf.read_bytes()
    assert copias[0].name.endswith(f"_{pdf.name}")
    # pasta = <ambiente>/<AAAA>/<MM> — em teste o ambiente ativo tem slug "test"
    assert copias[0].parent.parent.parent == _recebidos_root() / "test"


def test_commit_grava_sha256_e_caminho_da_copia_no_import(tmp_path, monkeypatch):
    import hashlib

    _config_em_tmp(monkeypatch, tmp_path)
    pdf = _sample_pdf()
    r = client.post(
        "/api/preview", files=[("file", (pdf.name, pdf.read_bytes(), "application/pdf"))]
    )
    preview_id = r.json()["preview_id"]

    r2 = client.post("/api/commit", json={"preview_id": preview_id})

    assert r2.status_code == 200, r2.text
    entry = client.get(f"/api/imported/{r2.json()['entry_id']}").json()["entry"]
    assert entry["file_sha256"] == hashlib.sha256(pdf.read_bytes()).hexdigest()
    assert Path(entry["original_path"]) == _arquivos_recebidos()[0]


def test_commit_grava_o_check_preco_da_fonte(tmp_path, monkeypatch):
    from app.erp import preco_trava
    from app.persistence import repo

    _config_em_tmp(monkeypatch, tmp_path)
    chamadas = []

    def fake(order, env):
        chamadas.append(order.header.order_number)
        return {"available": True, "items": [], "fonte": "mm"}

    monkeypatch.setattr(preco_trava, "check_preco_para_gravar", fake)
    pdf = _sample_pdf()
    r = client.post(
        "/api/preview", files=[("file", (pdf.name, pdf.read_bytes(), "application/pdf"))]
    )

    r2 = client.post("/api/commit", json={"preview_id": r.json()["preview_id"]})

    assert r2.status_code == 200, r2.text
    assert len(chamadas) == 1
    assert repo.get_import(r2.json()["entry_id"])["check_preco"]["fonte"] == "mm"


def test_preview_da_pasta_guarda_copia_e_o_move_para_importados_continua(tmp_path, monkeypatch):
    pdf = _sample_pdf()
    watch = tmp_path / "watch"
    watch.mkdir()
    seeded = watch / pdf.name
    seeded.write_bytes(pdf.read_bytes())
    imp = _config_em_tmp(monkeypatch, tmp_path, watch=watch)

    r = client.post("/api/preview-pending", json={"filename": pdf.name})
    assert r.status_code == 200, r.text
    assert len(_arquivos_recebidos()) == 1  # guardou no recebimento, antes do commit

    r2 = client.post("/api/commit", json={"preview_id": r.json()["preview_id"]})

    assert r2.status_code == 200, r2.text
    assert not seeded.exists() and (imp / pdf.name).exists()  # comportamento antigo intacto
    entry = client.get(f"/api/imported/{r2.json()['entry_id']}").json()["entry"]
    assert Path(entry["original_path"]) == _arquivos_recebidos()[0]  # aponta pra NOSSA cópia
    assert _arquivos_recebidos()[0].read_bytes() == pdf.read_bytes()


def test_arquivo_que_o_parser_nao_reconhece_ainda_e_guardado(monkeypatch):
    import app.pipeline

    monkeypatch.setattr(app.pipeline, "process_detailed", lambda _loaded, diag=None: None)

    r = client.post(
        "/api/preview", files=[("file", ("misterio.pdf", b"%PDF-1.4 lixo", "application/pdf"))]
    )

    assert r.status_code == 422
    copias = _arquivos_recebidos()
    assert len(copias) == 1
    assert copias[0].read_bytes() == b"%PDF-1.4 lixo"


def test_falha_ao_guardar_original_bloqueia_o_preview_com_erro_claro(monkeypatch):
    from app.ingestion import arquivo_recebido

    def _explode(*_a, **_k):
        raise OSError("disco cheio")

    monkeypatch.setattr(arquivo_recebido, "guardar", _explode)
    pdf = _sample_pdf()

    r = client.post(
        "/api/preview", files=[("file", (pdf.name, pdf.read_bytes(), "application/pdf"))]
    )

    assert r.status_code == 500
    assert "guardar o arquivo original" in r.json()["detail"]


def test_import_em_lote_da_pasta_guarda_copia_e_grava_sha(tmp_path, monkeypatch):
    import hashlib

    pdf = _sample_pdf()
    watch = tmp_path / "watch"
    watch.mkdir()
    (watch / pdf.name).write_bytes(pdf.read_bytes())
    _config_em_tmp(monkeypatch, tmp_path, watch=watch)

    r = client.post("/api/import", json={"files": [pdf.name]})

    assert r.status_code == 200, r.text
    assert r.json()["results"], r.text
    assert len(_arquivos_recebidos()) == 1
    entries = client.get("/api/imported?limit=5").json()["entries"]
    entry = client.get(f"/api/imported/{entries[0]['id']}").json()["entry"]
    assert entry["file_sha256"] == hashlib.sha256(pdf.read_bytes()).hexdigest()
    assert Path(entry["original_path"]) == _arquivos_recebidos()[0]


def test_download_do_arquivo_original_devolve_a_copia_guardada(tmp_path, monkeypatch):
    _config_em_tmp(monkeypatch, tmp_path)
    pdf = _sample_pdf()
    r = client.post(
        "/api/preview", files=[("file", (pdf.name, pdf.read_bytes(), "application/pdf"))]
    )
    entry_id = client.post("/api/commit", json={"preview_id": r.json()["preview_id"]}).json()[
        "entry_id"
    ]

    r2 = client.get(f"/api/imported/{entry_id}/arquivo-original")

    assert r2.status_code == 200, r2.text
    assert r2.content == pdf.read_bytes()
    assert pdf.name in r2.headers["content-disposition"]
    assert r2.headers["content-type"].startswith("application/pdf")


def test_download_do_arquivo_original_404_para_pedido_sem_copia():
    entry_id = _insert_import_minimo()

    r = client.get(f"/api/imported/{entry_id}/arquivo-original")

    assert r.status_code == 404
    assert "anterior" in r.json()["detail"]  # pedido anterior à guarda, mensagem amigável


def test_download_do_arquivo_original_recusa_caminho_fora_da_pasta_recebidos(tmp_path):
    fora = tmp_path / "fora.pdf"
    fora.write_bytes(b"segredo")
    entry_id = _insert_import_minimo(original_path=str(fora), file_sha256="x")

    r = client.get(f"/api/imported/{entry_id}/arquivo-original")

    assert r.status_code == 403


def test_rehydrate_preview_informa_se_ha_arquivo_original(tmp_path, monkeypatch):
    _config_em_tmp(monkeypatch, tmp_path)
    pdf = _sample_pdf()
    r = client.post(
        "/api/preview", files=[("file", (pdf.name, pdf.read_bytes(), "application/pdf"))]
    )
    entry_id = client.post("/api/commit", json={"preview_id": r.json()["preview_id"]}).json()[
        "entry_id"
    ]
    antigo = _insert_import_minimo(source_filename="antigo.pdf")

    com = client.get(f"/api/imported/{entry_id}/preview").json()["arquivo_original"]
    sem = client.get(f"/api/imported/{antigo}/preview")

    assert com == {"disponivel": True, "nome": pdf.name}
    # pedido antigo sem snapshot dá 422 no rehydrate; o que importa é o contrato do campo
    assert sem.status_code == 422


def test_send_to_fire_item_vinculado_vai_pelo_seq_do_depara(monkeypatch):
    """BACKLOG §2.1: no insert direto, item casado por de-para entra pelo SEQ."""
    import uuid
    from datetime import datetime

    from app.exporters import firebird_exporter as fb_mod
    from app.persistence import produto_depara_repo, repo

    entry_id = str(uuid.uuid4())
    repo.insert_import(
        {
            "id": entry_id,
            "source_filename": "pedido.pdf",
            "imported_at": datetime.now().isoformat(timespec="seconds"),
            "order_number": "TEST-DEP",
            "customer": "ACME",
            "status": "success",
            "portal_status": "parsed",
            "snapshot": {
                "header": {
                    "order_number": "TEST-DEP",
                    "customer_name": "ACME",
                    "customer_cnpj": "11111111000100",
                },
                "items": [
                    {"description": "a", "quantity": 1.0, "product_code": "REF-X"},
                    {"description": "b", "quantity": 1.0, "product_code": "SEM-VINCULO"},
                ],
                "source_file": "",
            },
        }
    )
    with db.connect() as conn:
        produto_depara_repo.upsert(
            conn,
            client_key=produto_depara_repo.client_key("11111111000100", None),
            chave_tipo="codigo",
            chave_valor="REF-X",
            fire_produto_id="77",
            fire_codigo="77",
            fire_ean=None,
            fire_nome="KIT",
            criado_em="t",
            criado_por="t",
        )

    recebido: dict = {}

    def _fake_export(self, order, *, override_client_id=None, produtos_por_item=None):
        recebido["produtos_por_item"] = produtos_por_item
        return fb_mod.FirebirdExportResult(
            order_number=order.header.order_number, items_inserted=2, fire_codigo=4242
        )

    monkeypatch.setattr(fb_mod.FirebirdExporter, "export", _fake_export)
    monkeypatch.setattr("app.erp.product_check.check_order", lambda order, env=None: {})
    monkeypatch.setattr(
        "app.erp.product_check.is_blocking", lambda check, ack_items=None: (False, {})
    )
    from app import config as app_config

    monkeypatch.setattr(
        app_config,
        "load",
        lambda: {"watch_dir": ".", "output_dir": ".", "export_mode": "db"},
    )

    r = client.post(f"/api/imported/{entry_id}/send-to-fire")
    assert r.status_code == 200, r.text
    assert recebido["produtos_por_item"] == {0: 77}


# ── Caminhos legados nunca gravam no Fire sem o preview ──────────────────────


@pytest.mark.parametrize("modo", ["db", "both"])
def test_import_direto_em_modo_db_nunca_chama_firebird(monkeypatch, tmp_path, modo):
    from app import config as app_config
    from app.exporters import firebird_exporter as fb_mod
    from app.web import server as server_mod

    def _explode(self, *a, **kw):
        raise AssertionError("FirebirdExporter.export não pode ser chamado fora do preview")

    monkeypatch.setattr(fb_mod.FirebirdExporter, "export", _explode)
    monkeypatch.setattr(
        app_config,
        "load",
        lambda: {"watch_dir": str(tmp_path), "output_dir": str(tmp_path), "export_mode": modo},
    )

    pdf = SAMPLES / "2600009562-2026-02-25.pdf"
    if not pdf.exists():
        pytest.skip("Sample PDF not available")
    from app.ingestion.file_loader import LoadedFile
    from app.pipeline import process

    loaded = LoadedFile(path=pdf, extension=".pdf", raw=pdf.read_bytes())
    order = process(loaded)

    out = server_mod._run_exporters(order, tmp_path)

    assert out["db_result"] == {"skipped": True, "skip_reason": "USE_PREVIEW"}
    assert out["fire_codigo"] is None
    if modo == "both":
        assert out["output_files"], "em both o XLSX ainda precisa ser gerado"
    else:
        assert out["output_files"] == []


def _import_parsed(entry_id, *, source_filename="x.pdf", file_sha256=None):
    from datetime import datetime

    from app.persistence import repo

    repo.insert_import(
        {
            "id": entry_id,
            "source_filename": source_filename,
            "file_sha256": file_sha256,
            "imported_at": datetime.now().isoformat(timespec="seconds"),
            "order_number": "TEST-PRECO",
            "customer": "ACME",
            "status": "success",
            "portal_status": "parsed",
            "snapshot": {"header": {"order_number": "TEST-PRECO"}, "items": []},
        }
    )


def _revisao_enviada(entry_id):
    from datetime import UTC, datetime

    from app.persistence import preco_revisao_repo as revs

    with db.connect() as conn:
        revs.garantir_revisao(conn, entry_id, {"externalId": entry_id}, "sha256:a", agora=datetime.now(UTC))
        revs.marcar_enviada(conn, entry_id, 1, {}, em=datetime.now(UTC))


def test_cancel_grava_retirada_da_revisao_enviada():
    import uuid

    from app.persistence import preco_revisao_repo as revs

    entry_id = str(uuid.uuid4())
    _import_parsed(entry_id)
    _revisao_enviada(entry_id)

    r = client.post(f"/api/imported/{entry_id}/cancel", json={"reason": "duplicado"})
    assert r.status_code == 200

    with db.connect() as conn:
        assert revs.vigente(conn, entry_id) is None
        pend = revs.retiradas_pendentes(conn, 10)
    assert [(p["import_id"], p["revisao"], p["retirada_motivo"]) for p in pend] == [
        (entry_id, 1, "duplicado")
    ]


def test_reimport_retira_a_revisao_do_pedido_antigo(monkeypatch, tmp_path):
    import shutil
    import uuid

    from app import config as app_config
    from app.persistence import preco_revisao_repo as revs

    pdf = SAMPLES / "2600009562-2026-02-25.pdf"
    if not pdf.exists():
        pytest.skip("Sample PDF not available")
    env = _env_trava_ligada(tmp_path)
    monkeypatch.setattr(
        app_config,
        "load",
        lambda: {"watch_dir": str(tmp_path), "output_dir": str(tmp_path), "export_mode": "xlsx"},
    )
    imp = tmp_path / "in" / "Pedidos importados"
    imp.mkdir(parents=True)
    shutil.copy(pdf, imp / pdf.name)
    import hashlib

    sha = hashlib.sha256(pdf.read_bytes()).hexdigest()
    # mesmos bytes com o nome carimbado de quando foi importado → é o antigo
    antigo = str(uuid.uuid4())
    _import_parsed(antigo, source_filename=f"20261017-1015_{pdf.name}", file_sha256=sha)
    _revisao_enviada(antigo)
    # homônimo com outro conteúdo → NÃO é o antigo
    outro = str(uuid.uuid4())
    _import_parsed(outro, source_filename=pdf.name, file_sha256="0" * 64)
    _revisao_enviada(outro)

    client.cookies.set("portal_env", env["id"])
    try:
        r = client.post("/api/reimport", json={"filename": pdf.name})
    finally:
        client.cookies.clear()

    assert r.status_code == 200, r.text
    with db.connect() as conn:
        assert revs.vigente(conn, antigo) is None
        assert revs.vigente(conn, outro) is not None
        pend = revs.retiradas_pendentes(conn, 10)
    assert [(p["import_id"], p["retirada_motivo"]) for p in pend] == [
        (antigo, "pedido reimportado no Portal")
    ]


def test_reimport_em_cadeia_retira_o_reimport_anterior(monkeypatch, tmp_path):
    import hashlib
    import shutil
    import uuid

    from app import config as app_config
    from app.persistence import preco_revisao_repo as revs

    pdf = SAMPLES / "2600009562-2026-02-25.pdf"
    if not pdf.exists():
        pytest.skip("Sample PDF not available")
    env = _env_trava_ligada(tmp_path)
    monkeypatch.setattr(
        app_config,
        "load",
        lambda: {"watch_dir": str(tmp_path), "output_dir": str(tmp_path), "export_mode": "xlsx"},
    )
    imp = tmp_path / "in" / "Pedidos importados"
    imp.mkdir(parents=True)
    shutil.copy(pdf, imp / pdf.name)
    original = str(uuid.uuid4())
    sha = hashlib.sha256(pdf.read_bytes()).hexdigest()
    _import_parsed(original, source_filename=f"20261017-1015_{pdf.name}", file_sha256=sha)

    client.cookies.set("portal_env", env["id"])
    try:
        r1 = client.post("/api/reimport", json={"filename": pdf.name})
        assert r1.status_code == 200, r1.text
        reimport_1 = r1.json()["entry_id"]
        _revisao_enviada(reimport_1)  # sem file_sha256: casa pelo nome guardado
        r2 = client.post("/api/reimport", json={"filename": pdf.name})
    finally:
        client.cookies.clear()

    assert r2.status_code == 200, r2.text
    reimport_2 = r2.json()["entry_id"]
    with db.connect() as conn:
        assert revs.vigente(conn, reimport_1) is None
        pend = revs.retiradas_pendentes(conn, 10)
    assert [p["import_id"] for p in pend] == [reimport_1]
    assert reimport_2 not in {p["import_id"] for p in pend}


# ── Trava de preço (preco_trava.decidir_envio) nos dois caminhos de envio ─────


_EMPRESA_TRAVA = "11111111000191"  # CNPJ do ambiente "test" = empresa faturadora
_R1_TRAVA = "a1a1a1a1-0000-4000-8000-000000000001"


def _pagina_trava(seq: str, centavos: int, cursor: str) -> None:
    """Página do contrato precos v1: rede r1 (raiz 11222333) com o produto `seq`."""
    import uuid
    from datetime import UTC
    from datetime import datetime as _dt

    from app.persistence import precos_repo

    with db.connect() as conn:
        precos_repo.aplicar_pagina(
            conn,
            {
                "schema": "precos.tabela.v1",
                "capacidades": ["precos.tabela.v1", "precos.aprovacao.v1"],
                "geradoEm": _dt.now(UTC).isoformat(),
                "empresaFaturadoraCnpj": _EMPRESA_TRAVA,
                "reiniciar": False,
                "mudancas": [
                    {
                        "tipo": "rede",
                        "redeId": _R1_TRAVA,
                        "nome": "Rede X",
                        "ativa": True,
                        "travada": False,
                        "raizes": ["11222333"],
                        "cnpjs": [],
                    },
                    {
                        "tipo": "preco",
                        "precoId": str(uuid.uuid5(uuid.NAMESPACE_URL, f"r1-{seq}")),
                        "redeId": _R1_TRAVA,
                        "empresaFaturadoraCnpj": _EMPRESA_TRAVA,
                        "erpProdutoId": seq,
                        "unidadeErp": "KIT",
                        "precoCentavos": centavos,
                        "vigenciaInicio": "2026-01-01",
                        "registradoEm": "2026-01-01T00:00:00.000Z",
                        "cancelado": False,
                    },
                ],
                "cursor": cursor,
                "temMais": False,
            },
            empresa_cnpj=_EMPRESA_TRAVA,
        )


def _env_trava_ligada(tmp_path):
    """Ambiente com slug "test" (mesmo SQLite do ambiente ativo dos testes),
    trava `ligado` e a cópia da tabela da rede 11222333 com o produto 123 a R$ 50,00."""
    from app.persistence import environments_repo

    env = environments_repo.create(
        slug="test",
        name="Teste",
        watch_dir=str(tmp_path / "in"),
        output_dir=str(tmp_path / "out"),
        fb_path=str(tmp_path / "x.fdb"),
        cnpj=_EMPRESA_TRAVA,
    )
    environments_repo.set_preco_trava_modo(env["id"], "ligado", por="teste")
    _pagina_trava("123", 5000, "c1")
    return env


def _seed_trava(entry_id, unit_price):
    from datetime import datetime as _dt

    from app.persistence import repo

    item = {"description": "x", "quantity": 10.0, "ean": "7891", "product_code": "C1"}
    if unit_price is not None:
        item["unit_price"] = unit_price
    repo.insert_import(
        {
            "id": entry_id,
            "source_filename": "p.pdf",
            "imported_at": _dt.now().isoformat(timespec="seconds"),
            "order_number": "TRAVA-1",
            "status": "success",
            "portal_status": "parsed",
            "snapshot": {
                "header": {
                    "order_number": "TRAVA-1",
                    "customer_name": "Rede X",
                    "customer_cnpj": "11.222.333/0001-44",
                },
                "items": [item],
                "source_file": "",
            },
        }
    )


def _check_trava(monkeypatch, seq=123):
    from app.erp import product_check as pc_mod

    monkeypatch.setattr(
        pc_mod,
        "check_order",
        lambda order, **kw: {
            "available": True,
            "items": [
                {
                    "ean": "7891",
                    "product_code": "C1",
                    "fire_product_id": seq,
                    "fire_unidade": "KIT",
                    "price_status": "mismatch",
                }
            ],
            "summary": {},
        },
    )


def test_send_to_fire_bloqueado_pela_trava_ligada(monkeypatch, tmp_path):
    import uuid

    from app.exporters import firebird_exporter as fb_mod
    from app.persistence import repo

    env = _env_trava_ligada(tmp_path)
    entry_id = str(uuid.uuid4())
    _seed_trava(entry_id, 54.0)  # 54 ≠ 50 da tabela → fila do Flow
    _check_trava(monkeypatch)
    monkeypatch.setattr(
        fb_mod.FirebirdExporter,
        "export",
        lambda *a, **kw: (_ for _ in ()).throw(AssertionError("não pode chamar o Fire")),
    )

    client.cookies.set("portal_env", env["id"])
    try:
        r = client.post(f"/api/imported/{entry_id}/send-to-fire")
    finally:
        client.cookies.clear()

    assert r.status_code == 409, r.text
    assert r.json()["detail"] == (
        "Pedido aguardando aprovação de preço no Flow: 1 item(ns)."
    )
    bloqueios = [a for a in repo.list_audit(entry_id) if a["event_type"] == "send_to_fire_blocked"]
    assert bloqueios[0]["detail"]["reason"] == "preco_aguardando"


def test_export_xlsx_leva_preco_aprovado(monkeypatch, tmp_path):
    import uuid

    from openpyxl import load_workbook

    env = _env_trava_ligada(tmp_path)
    entry_id = str(uuid.uuid4())
    _seed_trava(entry_id, None)  # pedido veio sem preço
    _check_trava(monkeypatch)

    client.cookies.set("portal_env", env["id"])
    try:
        r = client.post(f"/api/imported/{entry_id}/export-xlsx")
        assert r.status_code == 409, r.text
        did = _aprovar_preco_inicial(entry_id, 4800)
        r = client.post(f"/api/imported/{entry_id}/export-xlsx")
    finally:
        client.cookies.clear()

    assert r.status_code == 200, r.text
    path = r.json()["output_files"][0]["path"]
    ws = load_workbook(path).active
    valores = [c.value for row in ws.iter_rows() for c in row]
    assert 48.0 in valores
    from app.persistence import repo

    aplicados = [a["detail"] for a in repo.list_audit(entry_id) if a["event_type"] == "preco_aplicado"]
    assert aplicados == [{"item_index": 0, "preco": 48.0, "flow_id": did}]


def _decisao_flow(entry_id, resultado, centavos=None, motivo=None):
    """Decisão do Flow (rota c) para o item 0 na revisão vigente; devolve o decisaoId."""
    import uuid
    from datetime import UTC
    from datetime import datetime as _dt

    from app.persistence import preco_revisao_repo as revs

    did = str(uuid.uuid4())
    with db.connect() as conn:
        vig = revs.vigente(conn, entry_id)
        assert vig is not None, "o portão ainda não criou a revisão"
        assert revs.gravar_decisao(
            conn,
            {
                "decisaoId": did,
                "externalId": entry_id,
                "empresaFaturadoraCnpj": _EMPRESA_TRAVA,
                "erpProdutoId": "123",
                "revisao": vig["revisao"],
                "itemChave": f"{entry_id}:0",
                "resultado": resultado,
                "precoAprovadoCentavos": centavos,
                "motivoRecusa": motivo,
                "observacao": None,
                "decididoEm": "2026-10-18T12:00:00.000Z",
            },
            agora=_dt.now(UTC),
        ) == "aplicada"
    return did


def _aprovar_preco_inicial(entry_id, centavos):
    return _decisao_flow(entry_id, "acordo_comercial", centavos)


def test_send_to_fire_ligado_entrega_o_preco_aprovado_ao_exporter(monkeypatch, tmp_path):
    import uuid

    from app import config as app_config
    from app.exporters import firebird_exporter as fb_mod
    from app.persistence import repo

    env = _env_trava_ligada(tmp_path)
    entry_id = str(uuid.uuid4())
    _seed_trava(entry_id, None)
    _check_trava(monkeypatch)
    monkeypatch.setattr(
        app_config,
        "load",
        lambda: {"watch_dir": str(tmp_path), "output_dir": str(tmp_path), "export_mode": "db"},
    )
    recebidos: list = []

    def _export(self, order, **kw):
        recebidos.append([it.unit_price for it in order.items])
        return fb_mod.FirebirdExportResult(
            order_number=order.header.order_number, items_inserted=1, fire_codigo=777
        )

    monkeypatch.setattr(fb_mod.FirebirdExporter, "export", _export)

    client.cookies.set("portal_env", env["id"])
    try:
        r = client.post(f"/api/imported/{entry_id}/send-to-fire")
        assert r.status_code == 409, r.text
        assert recebidos == []
        assert not [a for a in repo.list_audit(entry_id) if a["event_type"] == "preco_aplicado"]
        _aprovar_preco_inicial(entry_id, 4800)
        r = client.post(f"/api/imported/{entry_id}/send-to-fire")
    finally:
        client.cookies.clear()

    assert r.status_code == 200, r.text
    assert recebidos == [[48.0]]
    aplicados = [a["detail"] for a in repo.list_audit(entry_id) if a["event_type"] == "preco_aplicado"]
    assert [(a["item_index"], a["preco"]) for a in aplicados] == [(0, 48.0)]


def test_send_to_fire_que_falha_no_fire_nao_grava_preco_aplicado(monkeypatch, tmp_path):
    import uuid

    from app import config as app_config
    from app.exporters import firebird_exporter as fb_mod
    from app.persistence import repo

    env = _env_trava_ligada(tmp_path)
    entry_id = str(uuid.uuid4())
    _seed_trava(entry_id, None)
    _check_trava(monkeypatch)
    monkeypatch.setattr(
        app_config,
        "load",
        lambda: {"watch_dir": str(tmp_path), "output_dir": str(tmp_path), "export_mode": "db"},
    )
    monkeypatch.setattr(
        fb_mod.FirebirdExporter,
        "export",
        lambda self, order, **kw: fb_mod.FirebirdExportResult(
            order_number=order.header.order_number,
            items_inserted=0,
            skipped=True,
            skip_reason="PRODUCT_NOT_FOUND",
        ),
    )

    client.cookies.set("portal_env", env["id"])
    try:
        assert client.post(f"/api/imported/{entry_id}/send-to-fire").status_code == 409
        _aprovar_preco_inicial(entry_id, 4800)  # liberado pela trava...
        r = client.post(f"/api/imported/{entry_id}/send-to-fire")
    finally:
        client.cookies.clear()

    assert r.status_code == 409, r.text  # ...mas o Fire recusou
    assert "não existe no Fire" in r.json()["detail"]
    assert not [a for a in repo.list_audit(entry_id) if a["event_type"] == "preco_aplicado"]


@pytest.mark.parametrize("modo", ["observando", "ligado"])
def test_run_exporters_com_trava_ativa_adia_o_xlsx(monkeypatch, tmp_path, modo):
    from app import config as app_config
    from app.models.order import Order, OrderHeader, OrderItem
    from app.web import server as server_mod

    monkeypatch.setattr(
        app_config,
        "load",
        lambda: {"watch_dir": str(tmp_path), "output_dir": str(tmp_path), "export_mode": "xlsx"},
    )
    order = Order(
        header=OrderHeader(order_number="A1"),
        items=[OrderItem(description="x", quantity=1, unit_price=10.0)],
    )

    out = server_mod._run_exporters(order, tmp_path, env={"slug": "mm", "preco_trava_modo": modo})

    assert out["output_files"] == []
    assert out["xlsx_adiado"] == "trava_de_preco"
    assert list(tmp_path.glob("*.xlsx")) == []


@pytest.mark.parametrize("env", [None, {"slug": "mm", "preco_trava_modo": "desligado"}])
def test_run_exporters_desligado_gera_xlsx_como_hoje(monkeypatch, tmp_path, env):
    from app import config as app_config
    from app.models.order import Order, OrderHeader, OrderItem
    from app.web import server as server_mod

    monkeypatch.setattr(
        app_config,
        "load",
        lambda: {"watch_dir": str(tmp_path), "output_dir": str(tmp_path), "export_mode": "xlsx"},
    )
    order = Order(
        header=OrderHeader(order_number="A1"),
        items=[OrderItem(description="x", quantity=1, unit_price=10.0)],
    )

    out = server_mod._run_exporters(order, tmp_path, env=env)

    assert len(out["output_files"]) == 1
    assert "xlsx_adiado" not in out
    assert set(out) == {"output_files", "db_result", "fire_codigo"}


def test_reimport_com_trava_ligada_nao_gera_xlsx_e_fica_em_revisao(monkeypatch, tmp_path):
    import shutil

    from app import config as app_config
    from app.persistence import repo

    pdf = SAMPLES / "2600009562-2026-02-25.pdf"
    if not pdf.exists():
        pytest.skip("Sample PDF not available")
    env = _env_trava_ligada(tmp_path)
    monkeypatch.setattr(
        app_config,
        "load",
        lambda: {"watch_dir": str(tmp_path), "output_dir": str(tmp_path), "export_mode": "xlsx"},
    )
    imp = tmp_path / "in" / "Pedidos importados"
    imp.mkdir(parents=True)
    shutil.copy(pdf, imp / pdf.name)

    client.cookies.set("portal_env", env["id"])
    try:
        r = client.post("/api/reimport", json={"filename": pdf.name})
    finally:
        client.cookies.clear()

    assert r.status_code == 200, r.text
    body = r.json()
    assert body["files"] == []
    assert body["xlsx_adiado"] == "trava_de_preco"
    assert list((tmp_path / "out").glob("*.xlsx")) == []
    assert repo.get_import(body["entry_id"])["portal_status"] == "parsed"


def test_override_cliente_grava_o_cnpj_do_fire(monkeypatch):
    from app.persistence import repo

    entry_id = _seed_parsed_entry()
    _patch_fb_with_rows(monkeypatch, [(4242, "ACME COMERCIO LTDA", "11.222.333/0001-44")])

    r = client.post(f"/api/imported/{entry_id}/override-cliente", json={"cliente_codigo": 4242})

    assert r.status_code == 200, r.text
    assert repo.get_import(entry_id)["cliente_override_cnpj"] == "11222333000144"


# ── Preview: situação de preço por item ──────────────────────────────────────


def _seed_preco_preview(entry_id, unit_price=54.0):
    import json
    from datetime import datetime as _dt

    from app.persistence import repo

    _seed_trava(entry_id, unit_price)
    with db.connect() as conn:
        conn.execute(
            "UPDATE imports SET check_json = ? WHERE id = ?",
            (
                json.dumps(
                    {
                        "available": True,
                        "items": [
                            {
                                "ean": "7891",
                                "product_code": "C1",
                                "fire_product_id": 123,
                                "fire_unidade": "KIT",
                                "price_status": "match",
                            }
                        ],
                        "summary": {},
                    }
                ),
                entry_id,
            ),
        )
    assert repo.get_import(entry_id)["check"]["items"], "check não gravado"
    return _dt.now()


def _preview_com_modo(entry_id, env, modo):
    from app.persistence import environments_repo

    environments_repo.set_preco_trava_modo(env["id"], modo, por="teste")
    client.cookies.set("portal_env", env["id"])
    try:
        return client.get(f"/api/imported/{entry_id}/preview")
    finally:
        client.cookies.clear()


def test_detalhe_traz_avaliacao_de_preco_em_observando(tmp_path):
    import uuid

    env = _env_trava_ligada(tmp_path)
    entry_id = str(uuid.uuid4())
    _seed_preco_preview(entry_id, 54.0)  # 54 ≠ 50 → fila

    r = _preview_com_modo(entry_id, env, "observando")

    assert r.status_code == 200, r.text
    preco = r.json()["preco"]
    assert preco["modo"] == "observando"
    assert preco["itens"] == [
        {
            "item_index": 0,
            "situacao": "fila",
            "motivo": "preco_acima",
            "preco_pedido": 54.0,
            "preco_tabela": 50.0,
            "decisao": None,
            "estado": "aguardando",
        }
    ]
    assert preco["bloqueio_pedido"] is None
    assert preco["resumo"] == {
        "ok": 0,
        "fila": 1,
        "aprovados": 0,
        "recusados": 0,
        "nao_avaliados": 0,
        "intercompany": 0,
    }


def test_detalhe_sem_bloco_preco_em_desligado(tmp_path):
    import uuid

    env = _env_trava_ligada(tmp_path)
    entry_id = str(uuid.uuid4())
    _seed_preco_preview(entry_id, 54.0)

    r = _preview_com_modo(entry_id, env, "desligado")

    assert r.status_code == 200, r.text
    assert "preco" not in r.json()


def _revisao_preview(entry_id):
    """Revisão vigente com o conteúdo atual do pedido (o que o portão criaria)."""
    from datetime import UTC
    from datetime import datetime as _dt

    from app.erp import preco_trava
    from app.models.order import Order
    from app.persistence import preco_revisao_repo as revs
    from app.persistence import repo

    entry = repo.get_import(entry_id)
    order = Order.model_validate(entry["snapshot"])
    env = {"preco_trava_modo": "ligado", "cnpj": _EMPRESA_TRAVA}
    agora = _dt.now(UTC)
    av = preco_trava.avaliar_entry(entry, order, env=env, agora=agora)
    h = preco_trava.conteudo_hash(order, av, cliente_cnpj=av.cnpj_rede, empresa_cnpj=_EMPRESA_TRAVA)
    corpo = {
        "externalId": entry_id,
        "empresaFaturadoraCnpj": _EMPRESA_TRAVA,
        "itens": [{"itemChave": f"{entry_id}:0", "erpProdutoId": "123"}],
    }
    with db.connect() as conn:
        revs.garantir_revisao(conn, entry_id, corpo, h, agora=agora)


def test_detalhe_mostra_decisao_do_item(tmp_path):
    import uuid

    env = _env_trava_ligada(tmp_path)
    entry_id = str(uuid.uuid4())
    _seed_preco_preview(entry_id, 54.0)
    _revisao_preview(entry_id)
    _decisao_flow(entry_id, "acordo_comercial", 5400)

    r = _preview_com_modo(entry_id, env, "ligado")

    preco = r.json()["preco"]
    assert preco["itens"][0]["decisao"] == "acordo_comercial"
    assert preco["itens"][0]["estado"] == "aprovado"
    assert preco["resumo"]["aprovados"] == 1
    assert preco["resumo"]["fila"] == 0
    assert preco["revisao"] == {"numero": 1, "situacao": "pendente_envio", "erro_flow": None}


def test_detalhe_sem_preco_quando_avaliacao_falha(tmp_path, monkeypatch):
    import uuid

    from app.erp import preco_trava

    env = _env_trava_ligada(tmp_path)
    entry_id = str(uuid.uuid4())
    _seed_preco_preview(entry_id, 54.0)
    monkeypatch.setattr(
        preco_trava, "avaliar_entry", lambda *a, **k: (_ for _ in ()).throw(RuntimeError("boom"))
    )

    r = _preview_com_modo(entry_id, env, "observando")

    assert r.status_code == 200, r.text
    assert "preco" not in r.json()


def _decidir_web(entry_id, resultado, motivo=None):
    _revisao_preview(entry_id)
    _decisao_flow(entry_id, resultado, None if resultado == "recusado" else 5400, motivo)


def test_detalhe_ligado_item_recusado(tmp_path):
    import uuid

    env = _env_trava_ligada(tmp_path)
    entry_id = str(uuid.uuid4())
    _seed_preco_preview(entry_id, 54.0)
    _decidir_web(entry_id, "recusado", "preço fora")

    preco = _preview_com_modo(entry_id, env, "ligado").json()["preco"]

    assert preco["itens"][0]["estado"] == "recusado"
    assert preco["itens"][0]["decisao"] == "recusado"
    assert preco["resumo"]["recusados"] == 1 and preco["resumo"]["fila"] == 0


def test_detalhe_ligado_item_nao_avaliado(tmp_path):
    import json
    import uuid

    env = _env_trava_ligada(tmp_path)
    entry_id = str(uuid.uuid4())
    _seed_preco_preview(entry_id, 54.0)
    with db.connect() as conn:
        conn.execute(
            "UPDATE imports SET check_json = ? WHERE id = ?",
            (json.dumps({"available": False, "reason": "fb", "items": []}), entry_id),
        )

    preco = _preview_com_modo(entry_id, env, "ligado").json()["preco"]

    assert preco["itens"][0]["estado"] == "nao_avaliado"
    assert preco["itens"][0]["motivo"] == "produto_nao_identificado"
    assert preco["resumo"]["nao_avaliados"] == 1


def test_detalhe_ligado_item_intercompany(tmp_path):
    import uuid

    from app.persistence import environments_repo

    env = _env_trava_ligada(tmp_path)
    outro = environments_repo.create(
        slug="grupo",
        name="Grupo",
        watch_dir=str(tmp_path / "gi"),
        output_dir=str(tmp_path / "go"),
        fb_path="",
    )
    environments_repo.set_intercompany_config(outro["id"], cnpj="11.222.333/0001-44", revenda_slug="grupo")
    entry_id = str(uuid.uuid4())
    _seed_preco_preview(entry_id, 54.0)

    preco = _preview_com_modo(entry_id, env, "ligado").json()["preco"]

    assert preco["itens"][0]["estado"] == "intercompany"
    assert preco["resumo"]["intercompany"] == 1


# ── Revisão final da trava de preço ──────────────────────────────────────────


def _vinculo_de_para(seq: str, codigo: str):
    """Vínculo de-para do EAN 7891 do cliente 11222333000144 para o SEQ `seq`."""
    from app.persistence import produto_depara_repo

    with db.connect() as conn:
        produto_depara_repo.upsert(
            conn,
            client_key="11222333000144",
            chave_tipo="ean",
            chave_valor="7891",
            fire_produto_id=seq,
            fire_codigo=codigo,
            fire_ean=None,
            fire_nome="Produto B",
            criado_em="2026-10-18T10:00:00",
            criado_por="teste",
        )


def _preco_na_tabela(seq: str, centavos: int):
    _pagina_trava(seq, centavos, "c2")


def _cfg_db(monkeypatch, tmp_path, modo="db"):
    from app import config as app_config

    monkeypatch.setattr(
        app_config,
        "load",
        lambda: {"watch_dir": str(tmp_path), "output_dir": str(tmp_path), "export_mode": modo},
    )


@pytest.mark.parametrize("rota", ["send-to-fire", "export-xlsx"])
def test_trava_avalia_o_produto_do_vinculo_e_bloqueia(monkeypatch, tmp_path, rota):
    """Item casa por EAN com o SEQ 123 (R$ 50 = preço do pedido), mas o vínculo
    de-para leva ao SEQ 456 (R$ 60): o Fire/XLSX recebe 456, a trava avalia 456.
    A unidade do check é a do 123: a do 456 é desconhecida e bloqueia."""
    import uuid

    from app.exporters import firebird_exporter as fb_mod

    env = _env_trava_ligada(tmp_path)
    _preco_na_tabela("456", 6000)
    _vinculo_de_para("456", "B456")
    entry_id = str(uuid.uuid4())
    _seed_trava(entry_id, 50.0)
    _check_trava(monkeypatch)
    _cfg_db(monkeypatch, tmp_path)
    monkeypatch.setattr(
        fb_mod.FirebirdExporter,
        "export",
        lambda *a, **kw: (_ for _ in ()).throw(AssertionError("não pode chamar o Fire")),
    )

    client.cookies.set("portal_env", env["id"])
    try:
        r = client.post(f"/api/imported/{entry_id}/{rota}")
    finally:
        client.cookies.clear()

    assert r.status_code == 409, r.text
    assert "unidade do produto no Fire" in r.json()["detail"]


def test_trava_libera_pelo_vinculo_e_o_exporter_recebe_o_mesmo_seq(monkeypatch, tmp_path):
    import uuid

    from app.exporters import firebird_exporter as fb_mod

    env = _env_trava_ligada(tmp_path)
    _preco_na_tabela("123", 9999)  # A: preço diferente do pedido
    _preco_na_tabela("456", 5000)  # B: igual ao pedido
    _vinculo_de_para("456", "B456")
    entry_id = str(uuid.uuid4())
    _seed_trava(entry_id, 50.0)
    _check_trava(monkeypatch, seq=456)  # check_order resolveu pelo degrau de-para
    _cfg_db(monkeypatch, tmp_path)
    recebidos: list = []

    def _export(self, order, **kw):
        recebidos.append(kw.get("produtos_por_item"))
        return fb_mod.FirebirdExportResult(
            order_number=order.header.order_number, items_inserted=1, fire_codigo=777
        )

    monkeypatch.setattr(fb_mod.FirebirdExporter, "export", _export)

    client.cookies.set("portal_env", env["id"])
    try:
        r = client.post(f"/api/imported/{entry_id}/send-to-fire")
    finally:
        client.cookies.clear()

    assert r.status_code == 200, r.text
    assert recebidos == [{0: 456}]


def test_registrar_aplicados_que_falha_nao_derruba_o_envio(monkeypatch, tmp_path):
    import uuid

    from app.erp import preco_trava
    from app.exporters import firebird_exporter as fb_mod
    from app.persistence import repo
    from app.web import server

    env = _env_trava_ligada(tmp_path)
    entry_id = str(uuid.uuid4())
    _seed_trava(entry_id, None)
    _check_trava(monkeypatch)
    _cfg_db(monkeypatch, tmp_path)
    monkeypatch.setattr(
        fb_mod.FirebirdExporter,
        "export",
        lambda self, order, **kw: fb_mod.FirebirdExportResult(
            order_number=order.header.order_number, items_inserted=1, fire_codigo=777
        ),
    )
    pushes: list = []
    monkeypatch.setattr(server, "push_new_order", lambda *a, **kw: pushes.append(kw) or True)

    def _boom(*a, **kw):
        raise RuntimeError("audit caiu")

    client.cookies.set("portal_env", env["id"])
    try:
        assert client.post(f"/api/imported/{entry_id}/send-to-fire").status_code == 409
        _aprovar_preco_inicial(entry_id, 4800)
        monkeypatch.setattr(preco_trava, "registrar_aplicados", _boom)
        r = client.post(f"/api/imported/{entry_id}/send-to-fire")
    finally:
        client.cookies.clear()

    assert r.status_code == 200, r.text
    assert repo.get_import(entry_id)["portal_status"] == "sent_to_fire"
    assert len(pushes) == 1


def test_registrar_aplicados_que_falha_nao_pula_o_push_do_xlsx(monkeypatch, tmp_path):
    import uuid

    from app.erp import preco_trava
    from app.web import server

    env = _env_trava_ligada(tmp_path)
    entry_id = str(uuid.uuid4())
    _seed_trava(entry_id, None)
    _check_trava(monkeypatch)
    _cfg_db(monkeypatch, tmp_path, modo="xlsx")
    pushes: list = []
    monkeypatch.setattr(server, "push_new_order", lambda *a, **kw: pushes.append(kw) or True)

    def _boom(*a, **kw):
        raise RuntimeError("audit caiu")

    client.cookies.set("portal_env", env["id"])
    try:
        assert client.post(f"/api/imported/{entry_id}/export-xlsx").status_code == 409
        _aprovar_preco_inicial(entry_id, 4800)
        monkeypatch.setattr(preco_trava, "registrar_aplicados", _boom)
        r = client.post(f"/api/imported/{entry_id}/export-xlsx")
    finally:
        client.cookies.clear()

    assert r.status_code == 200, r.text
    assert len(pushes) == 1


@pytest.mark.parametrize("modo", ["observando", "ligado"])
def test_api_process_recusa_com_trava_ativa(tmp_path, modo):
    from app.persistence import environments_repo

    env = _env_trava_ligada(tmp_path)
    environments_repo.set_preco_trava_modo(env["id"], modo, por="teste")

    client.cookies.set("portal_env", env["id"])
    try:
        r = client.post(
            "/api/process",
            data={"output_dir": str(tmp_path)},
            files=[("files", ("p.pdf", b"%PDF-1.4", "application/pdf"))],
        )
    finally:
        client.cookies.clear()

    assert r.status_code == 409, r.text
    assert r.json()["detail"] == (
        "Com a trava de preço ativa, use a revisão do pedido para exportar."
    )
    assert list(tmp_path.glob("*.xlsx")) == []


def test_api_process_desligado_segue_como_hoje(tmp_path):
    from app.persistence import environments_repo

    env = _env_trava_ligada(tmp_path)
    environments_repo.set_preco_trava_modo(env["id"], "desligado", por="teste")

    client.cookies.set("portal_env", env["id"])
    try:
        r = client.post(
            "/api/process",
            data={"output_dir": str(tmp_path)},
            files=[("files", ("malware.exe", b"MZ", "application/octet-stream"))],
        )
    finally:
        client.cookies.clear()

    assert r.status_code == 200
    assert any("suportado" in e["error"].lower() for e in r.json()["errors"])
