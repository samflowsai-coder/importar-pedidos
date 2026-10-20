"""Saída para o Fire com o produto decidido (spec §6, Task 6).

Em `produto_chave_modo` `ligado`, o produto que a trava aprovou é o produto gravado:
o insert entra pelo SEQ de todo item `ok` e a planilha leva o `CODPROD_ALTERN` dele,
com o código do cliente na OBS. A planilha trava quando o importador do Fire (que
casa por PREFIXO, BACKLOG 2.15) trocaria o produto, ou quando não há código.
"""

from __future__ import annotations

import uuid
from datetime import UTC, datetime
from unittest.mock import patch

import pytest

from app.erp import depara_apply, preco_trava, product_check, produto_trava
from app.exporters.erp_exporter import ERPExporter
from app.models.order import Order, OrderHeader, OrderItem
from app.persistence import db
from tests.test_product_check import _fire_fisico

LIGADO = {"slug": "mm", "produto_chave_modo": "ligado"}


@pytest.fixture(autouse=True)
def isolated_sqlite(tmp_path_factory):
    tmp = tmp_path_factory.mktemp("dbstate")
    db.set_db_path(tmp / "app_state.db")
    db.reset_init_cache()
    yield
    db.set_db_path(None)
    db.reset_init_cache()


def _order(*itens: dict) -> Order:
    return Order(
        header=OrderHeader(order_number="P1", customer_cnpj="00000000000100", customer_name="ACME"),
        items=[OrderItem(description="x", quantity=1.0, unit_price=1.0, **kw) for kw in itens],
    )


def _check_fire(produtos, order: Order, **kw) -> tuple[dict, dict]:
    with patch("app.erp.product_check.FirebirdConnection") as mock_fb:
        mock_fb.return_value.is_configured.return_value = True
        ctx, estado = _fire_fisico(produtos, **kw)
        mock_fb.return_value.connect.return_value = ctx
        return product_check.check_order(order), estado


def _ok(seq, codprod, *, codigo="C1", troca=None, unidade="PAR", ean=None, **extra):
    produto = {
        "estado": "ok",
        "fire_seq": seq,
        "fonte": "ean",
        "motivo": "",
        "candidatos": [
            {
                "seq": seq,
                "descricao": "D",
                "codprod_altern": codprod,
                "bloqueado": False,
                "unidade": unidade,
                "ean": ean,
            }
        ],
        "troca_no_fire": troca,
    }
    produto.update(extra)
    return {"product_code": codigo, "ean": None, "produto": produto}


def _check(*itens) -> dict:
    return {"available": True, "produto_conferencia_incompleta": False, "items": list(itens)}


# ── check: previsão do importador sobre o código decidido ───────────────────


def test_check_preve_o_codigo_decidido_numa_passada_so():
    """Item casado por EAN com código do cliente diferente: a previsão roda sobre
    o código do cliente (aviso legado) e sobre o CODPROD_ALTERN decidido, numa query."""
    produtos = [(2077, "KG02PR", "KIT GRAZI", 1.0, "7891234567890")]
    order = _order({"product_code": "REF-9", "ean": "7891234567890"})
    check, estado = _check_fire(produtos, order)
    produto = check["items"][0]["produto"]
    assert produto["estado"] == "ok" and produto["fire_seq"] == 2077
    assert produto["troca_no_fire"] is None
    assert len(estado["previsoes"]) == 1
    assert sorted(estado["previsoes"][0]) == ["KG02PR", "REF-9"]


def test_check_marca_troca_quando_o_prefixo_vem_antes():
    produtos = [(3103, "NB01-3GG", "KIT NB01 GG", 1.0), (3102, "NB01-3G", "KIT NB01 G", 1.0)]
    check, _ = _check_fire(produtos, _order({"product_code": "NB01-3G"}))
    produto = check["items"][0]["produto"]
    assert produto["fire_seq"] == 3102
    assert produto["troca_no_fire"] == {"seq_previsto": 3103, "codigo": "NB01-3G"}


def test_check_sem_troca_quando_o_decidido_vem_primeiro():
    produtos = [(3102, "NB01-3G", "KIT NB01 G", 1.0), (3103, "NB01-3GG", "KIT NB01 GG", 1.0)]
    check, _ = _check_fire(produtos, _order({"product_code": "NB01-3G"}))
    assert check["items"][0]["produto"]["troca_no_fire"] is None


def test_check_previsao_falhou_marca_que_nao_da_para_conferir():
    produtos = [(3102, "NB01-3G", "KIT NB01 G", 1.0)]
    check, _ = _check_fire(produtos, _order({"product_code": "NB01-3G"}), falha_previsao=True)
    assert check["items"][0]["produto"]["troca_no_fire"] == {
        "seq_previsto": None,
        "codigo": "NB01-3G",
    }
    assert check["items"][0]["troca_no_fire"] is None  # o aviso legado não muda


def test_check_item_fora_de_ok_nao_tem_troca():
    check, _ = _check_fire([(1, "KG07BR", "KIT", 1.0)], _order({"product_code": "KG07"}))
    produto = check["items"][0]["produto"]
    assert produto["estado"] == "a_identificar"
    assert produto["troca_no_fire"] is None


def test_check_indisponivel_traz_a_chave_de_troca():
    with patch("app.erp.product_check.FirebirdConnection") as mock_fb:
        mock_fb.return_value.is_configured.return_value = False
        check = product_check.check_order(_order({"product_code": "X"}))
    assert check["items"][0]["produto"]["troca_no_fire"] is None


# ── trava: regras da planilha ────────────────────────────────────────────────


def test_planilha_trava_quando_previsao_de_prefixo_diverge():
    produtos = [(3103, "NB01-3GG", "KIT NB01 GG", 1.0), (3102, "NB01-3G", "KIT NB01 G", 1.0)]
    check, _ = _check_fire(produtos, _order({"product_code": "NB01-3G"}))
    r = produto_trava.avaliar(check, env=LIGADO, export_mode="xlsx")
    assert (r.bloqueado, r.reason) == (True, "troca_no_fire")
    assert r.itens == [
        {
            "index": 0,
            "estado": "troca_no_fire",
            "motivo": "O importador do Fire trocaria este código pelo produto 3103.",
        }
    ]
    # O insert entra pelo SEQ: a previsão da planilha não vale.
    assert produto_trava.avaliar(check, env=LIGADO, export_mode="db").bloqueado is False
    assert produto_trava.avaliar(check, env=LIGADO, export_mode="both").reason == "troca_no_fire"


def test_planilha_trava_sem_codprod_altern():
    r = produto_trava.avaliar(_check(_ok(2077, None)), env=LIGADO, export_mode="xlsx")
    assert (r.bloqueado, r.reason) == (True, "sem_codigo_fire")
    assert r.itens == [
        {
            "index": 0,
            "estado": "sem_codigo_fire",
            "motivo": "O produto 2077 não tem código no Fire; a planilha não consegue apontá-lo.",
        }
    ]
    assert produto_trava.avaliar(
        _check(_ok(2077, "  ")), env=LIGADO, export_mode="xlsx"
    ).reason == ("sem_codigo_fire")
    assert (
        produto_trava.avaliar(_check(_ok(2077, None)), env=LIGADO, export_mode="db").bloqueado
        is False
    )


def test_planilha_previsao_impossivel_trava_como_troca():
    troca = {"seq_previsto": None, "codigo": "KG02PR"}
    r = produto_trava.avaliar(
        _check(_ok(2077, "KG02PR", troca=troca)), env=LIGADO, export_mode="xlsx"
    )
    assert (r.bloqueado, r.reason) == (True, "troca_no_fire")
    assert r.itens[0]["motivo"] == (
        "O código KG02PR do produto 2077 não pode ser conferido na planilha do Fire."
    )


def test_planilha_check_sem_a_chave_de_troca_nao_libera():
    item = _ok(2077, "KG02PR")
    del item["produto"]["troca_no_fire"]
    r = produto_trava.avaliar(_check(item), env=LIGADO, export_mode="xlsx")
    assert r.reason == "troca_no_fire"


def test_planilha_precedencia_dos_motivos():
    troca = {"seq_previsto": 9, "codigo": "B"}
    pendente = {"product_code": "Z", "produto": {"estado": "a_identificar", "motivo": "m"}}
    ok_troca = _ok(2, "B", troca=troca)
    sem_codigo = _ok(3, None)
    r = produto_trava.avaliar(_check(ok_troca, sem_codigo), env=LIGADO, export_mode="xlsx")
    assert r.reason == "sem_codigo_fire" and [i["index"] for i in r.itens] == [1]
    r = produto_trava.avaliar(
        _check(ok_troca, sem_codigo, pendente), env=LIGADO, export_mode="xlsx"
    )
    assert r.reason == "produto_a_identificar"
    indisponivel = {**_check(ok_troca), "available": False}
    assert produto_trava.avaliar(indisponivel, env=LIGADO, export_mode="xlsx").reason == (
        "fire_indisponivel"
    )


def test_planilha_tudo_certo_libera():
    r = produto_trava.avaliar(_check(_ok(1, "A"), _ok(2, "B")), env=LIGADO, export_mode="xlsx")
    assert r == produto_trava.TravaProduto(bloqueado=False, reason=None, itens=[])


def test_observando_lista_as_regras_da_planilha_sem_travar():
    env = {"produto_chave_modo": "observando"}
    r = produto_trava.avaliar(_check(_ok(3, None)), env=env, export_mode="xlsx")
    assert r.bloqueado is False
    assert [i["estado"] for i in r.itens] == ["sem_codigo_fire"]


def test_desligado_ignora_as_regras_da_planilha():
    r = produto_trava.avaliar(_check(_ok(3, None)), env=None, export_mode="xlsx")
    assert r == produto_trava.TravaProduto(bloqueado=False, reason=None, itens=[])


# ── produtos_decididos → planilha e insert ──────────────────────────────────


def test_produtos_decididos_cobre_todo_item_ok():
    pendente = {"product_code": "Z", "produto": {"estado": "ambiguo", "fire_seq": None}}
    check = _check(
        _ok(10, "A", codigo="CLI-A", ean="7891234567890"), pendente, _ok(30, "C", codigo="C")
    )
    vinculos, mapa = produto_trava.produtos_decididos(check)
    assert mapa == {0: 10, 2: 30}
    comum = {"substitui_ean": True}
    assert vinculos == [
        {
            "item_index": 0,
            "fire_produto_id": 10,
            "fire_codigo": "A",
            "fire_ean": "7891234567890",
            "orig": "CLI-A",
        }
        | comum,
        {"item_index": 2, "fire_produto_id": 30, "fire_codigo": "C", "fire_ean": "", "orig": "C"}
        | comum,
    ]


def test_produtos_decididos_sem_check_e_vazio():
    assert produto_trava.produtos_decididos(None) == ([], {})
    assert produto_trava.produtos_decididos({"items": [{"descricao": "x"}]}) == ([], {})


def test_planilha_leva_codprod_altern_e_ref_cliente_na_obs():
    """Cliente manda KG02PR, casa pelo EAN no SEQ 2077 (CODPROD_ALTERN KG02PR)."""
    produtos = [(2077, "KG02PR", "KIT GRAZI", 1.0, "7891234567890")]
    order = _order({"product_code": "KG02PR", "ean": "7891234567890"})
    check, _ = _check_fire(produtos, order)
    assert check["items"][0]["produto"]["fonte"] == "ean"
    vinculos, mapa = produto_trava.produtos_decididos(check)
    depara_apply.aplicar(order, vinculos)
    (row,) = ERPExporter()._to_erp_rows(order, order.items)
    assert row.codigo_produto == "KG02PR"
    assert row.ean == "7891234567890"
    assert "ref cliente: KG02PR" in row.obs
    assert mapa == {0: 2077}


def test_planilha_troca_o_codigo_do_cliente_pelo_do_fire():
    produtos = [(2077, "KG02PR", "KIT GRAZI", 1.0, "7891234567890")]
    order = _order({"product_code": "REF-9", "ean": "7891234567890"})
    check, _ = _check_fire(produtos, order)
    depara_apply.aplicar(order, produto_trava.produtos_decididos(check)[0])
    (row,) = ERPExporter()._to_erp_rows(order, order.items)
    assert row.codigo_produto == "KG02PR"
    assert row.obs.endswith("ref cliente: REF-9")


# ── decidir_envio: a mesma decisão vai à avaliação de preço e ao chamador ───


def test_decidir_envio_ligado_devolve_o_check_aprovado(monkeypatch):
    check = _check(_ok(10, "A"))
    monkeypatch.setattr(product_check, "check_order", lambda order, **kw: check)
    monkeypatch.setattr(preco_trava, "_export_mode", lambda: "xlsx")
    d = preco_trava.decidir_envio(
        "id", _order({"product_code": "C1"}), {}, env=LIGADO, agora=datetime.now(UTC)
    )
    assert d.bloqueado is False
    assert d.check is check


@pytest.mark.parametrize("m", ["desligado", "observando"])
def test_decidir_envio_fora_de_ligado_nao_devolve_check(monkeypatch, m):
    monkeypatch.setattr(product_check, "check_order", lambda order, **kw: _check(_ok(10, "A")))
    env = {"slug": "mm", "produto_chave_modo": m}
    d = preco_trava.decidir_envio(
        "id", _order({"product_code": "C1"}), {}, env=env, agora=datetime.now(UTC)
    )
    assert d.check is None


def test_decidir_envio_ligado_avalia_preco_pelo_seq_decidido(monkeypatch):
    monkeypatch.setattr(product_check, "check_order", lambda order, **kw: _check(_ok(42, "A")))
    monkeypatch.setattr(preco_trava, "_export_mode", lambda: "xlsx")
    vistos = []

    def _ligado(import_id, order, entry, live, ack, agora, produtos_por_item, env):
        vistos.append(produtos_por_item)
        return preco_trava.DecisaoEnvio(False)

    monkeypatch.setattr(preco_trava, "_decidir_ligado", _ligado)
    env = {**LIGADO, "preco_trava_modo": "ligado"}
    d = preco_trava.decidir_envio(
        "id",
        _order({"product_code": "C1"}),
        {},
        env=env,
        agora=datetime.now(UTC),
        produtos_por_item={0: 999},
    )
    assert vistos == [{0: 42}]
    assert d.check is not None


def test_decidir_envio_export_mode_explicito_vale_para_a_trava(monkeypatch):
    """A rota de planilha passa "xlsx" mesmo com EXPORT_MODE=db: a planilha sai."""
    troca = {"seq_previsto": 9, "codigo": "A"}
    monkeypatch.setattr(
        product_check, "check_order", lambda o, **kw: _check(_ok(10, "A", troca=troca))
    )
    monkeypatch.setattr(preco_trava, "_export_mode", lambda: "db")
    order = _order({"product_code": "C1"})
    agora = datetime.now(UTC)
    assert preco_trava.decidir_envio("id", order, {}, env=LIGADO, agora=agora).bloqueado is False
    d = preco_trava.decidir_envio("id", order, {}, env=LIGADO, agora=agora, export_mode="xlsx")
    assert (d.bloqueado, d.reason) == (True, "troca_no_fire")
    assert d.detalhe == "Item 1: O importador do Fire trocaria este código pelo produto 9."


# ── rotas: o que sai é o que foi decidido ───────────────────────────────────


def _env_ligado(tmp_path):
    from app.persistence import environments_repo

    env = environments_repo.create(
        slug="test",
        name="Teste",
        watch_dir=str(tmp_path / "in"),
        output_dir=str(tmp_path / "out"),
        fb_path=str(tmp_path / "x.fdb"),
    )
    environments_repo.set_produto_chave_modo(env["id"], "ligado", por="teste")
    return environments_repo.get(env["id"])


def _seed(order: Order) -> str:
    from app.persistence import repo

    import_id = str(uuid.uuid4())
    repo.insert_import(
        {
            "id": import_id,
            "source_filename": "p.pdf",
            "imported_at": datetime.now().isoformat(timespec="seconds"),
            "order_number": order.header.order_number,
            "status": "success",
            "portal_status": "parsed",
            "snapshot": order.model_dump(),
        }
    )
    return import_id


def _config(monkeypatch, tmp_path, export_mode):
    from app import config as app_config

    monkeypatch.setattr(
        app_config,
        "load",
        lambda: {"output_dir": str(tmp_path / "out"), "export_mode": export_mode},
    )


def _sem_depara_v1(monkeypatch):
    """O de-para v1 apontaria outro SEQ: em `ligado` ele não pode sair."""
    from app.web import server as server_mod

    v1 = [
        {
            "item_index": 0,
            "fire_produto_id": 777,
            "fire_codigo": "777",
            "fire_ean": None,
            "orig": "x",
        }
    ]
    monkeypatch.setattr(server_mod, "_vinculos_de_para", lambda order, caminho, entry=None: (v1, {0: 777}))


def _nunca_v1(monkeypatch):
    """Em `ligado` o de-para v1 nem é lido."""
    from app.web import server as server_mod

    def _nunca(order, caminho):
        raise AssertionError(f"leu o de-para v1 em ligado ({caminho})")

    monkeypatch.setattr(server_mod, "_vinculos_de_para", _nunca)


def test_insert_direto_usa_seq_de_todo_item_ok(monkeypatch, tmp_path):
    from app.exporters import firebird_exporter
    from app.exporters.firebird_exporter import FirebirdExportResult
    from app.web import server as server_mod

    env = _env_ligado(tmp_path)
    _config(monkeypatch, tmp_path, "db")
    _nunca_v1(monkeypatch)
    order = _order({"product_code": "CLI-A"}, {"product_code": "B"})
    import_id = _seed(order)
    check = _check(_ok(10, "A", codigo="CLI-A"), _ok(20, None, codigo="B"))
    monkeypatch.setattr(product_check, "check_order", lambda o, **kw: check)
    vistos = []

    def _export(self, order, *, override_client_id=None, produtos_por_item=None):
        vistos.append((produtos_por_item, [it.product_code for it in order.items]))
        return FirebirdExportResult("P1", 0, skipped=True, skip_reason="TESTE")

    monkeypatch.setattr(firebird_exporter.FirebirdExporter, "export", _export)
    cfg = {"output_dir": str(tmp_path / "out"), "export_mode": "db"}
    out = server_mod._send_one_to_fire(import_id, cfg, request_env=env)
    assert out.reason == "TESTE"
    # Sem CODPROD_ALTERN, o insert entra pelo SEQ e o código do cliente fica.
    assert vistos == [({0: 10, 1: 20}, ["A", "B"])]


def test_export_xlsx_leva_o_codigo_decidido(monkeypatch, tmp_path):
    from app.web import server as server_mod

    env = _env_ligado(tmp_path)
    _config(monkeypatch, tmp_path, "xlsx")
    _nunca_v1(monkeypatch)
    order = _order({"product_code": "REF-9", "ean": "7891234567890"})
    import_id = _seed(order)
    monkeypatch.setattr(
        product_check, "check_order", lambda o, **kw: _check(_ok(2077, "KG02PR", codigo="REF-9"))
    )
    vistos = []
    monkeypatch.setattr(
        ERPExporter,
        "export",
        lambda self, order, out: vistos.append(order.model_copy(deep=True)) or [],
    )
    cfg = {"output_dir": str(tmp_path / "out"), "export_mode": "xlsx"}
    out = server_mod._export_one_xlsx(import_id, cfg, request_env=env)
    assert out.ok, out.detail
    (saiu,) = vistos
    assert saiu.items[0].product_code == "KG02PR"
    assert saiu.items[0].obs == "ref cliente: REF-9"
    # O produto decidido não tem EAN: o do cliente não vai na planilha.
    assert saiu.items[0].ean is None


def test_export_xlsx_com_export_mode_db_ainda_aplica_regra_da_planilha(monkeypatch, tmp_path):
    from app.web import server as server_mod

    env = _env_ligado(tmp_path)
    _config(monkeypatch, tmp_path, "db")
    import_id = _seed(_order({"product_code": "C1"}))
    monkeypatch.setattr(product_check, "check_order", lambda o, **kw: _check(_ok(5, None)))

    def _nunca(*a, **kw):
        raise AssertionError("gerou planilha sem código do Fire")

    monkeypatch.setattr(ERPExporter, "export", _nunca)
    cfg = {"output_dir": str(tmp_path / "out"), "export_mode": "db"}
    out = server_mod._export_one_xlsx(import_id, cfg, request_env=env)
    assert (out.ok, out.reason, out.http_status) == (False, "sem_codigo_fire", 409)


def test_export_xlsx_desligado_segue_o_de_para_de_hoje(monkeypatch, tmp_path):
    from app.persistence import environments_repo
    from app.web import server as server_mod

    env = _env_ligado(tmp_path)
    environments_repo.set_produto_chave_modo(env["id"], "desligado", por="teste")
    env = environments_repo.get(env["id"])
    _config(monkeypatch, tmp_path, "xlsx")
    _sem_depara_v1(monkeypatch)
    import_id = _seed(_order({"product_code": "C1"}))
    monkeypatch.setattr(
        product_check,
        "check_order",
        lambda o, **kw: {
            **_check(_ok(2077, "KG02PR", codigo="C1")),
            "items": [{**_ok(2077, "KG02PR", codigo="C1"), "price_status": "match"}],
        },
    )
    vistos = []
    monkeypatch.setattr(
        ERPExporter,
        "export",
        lambda self, order, out: vistos.append(order.model_copy(deep=True)) or [],
    )
    cfg = {"output_dir": str(tmp_path / "out"), "export_mode": "xlsx"}
    out = server_mod._export_one_xlsx(import_id, cfg, request_env=env)
    assert out.ok, out.detail
    assert vistos[0].items[0].product_code == "777"  # o vínculo v1, como hoje


# ── CLI ─────────────────────────────────────────────────────────────────────


def test_cli_ligado_exporta_o_codigo_decidido(monkeypatch, tmp_path):
    import main as cli

    monkeypatch.setenv("PORTAL_PRODUTO_CHAVE_CLI", "ligado")
    monkeypatch.setattr(
        cli.app_config, "load", lambda: {"export_mode": "xlsx", "output_dir": str(tmp_path)}
    )
    order = _order({"product_code": "REF-9"})
    monkeypatch.setattr(cli.FileLoader, "load_files", lambda self, d: [order])
    monkeypatch.setattr(cli, "process", lambda f: f)
    monkeypatch.setattr(
        cli, "check_order", lambda o, **kw: _check(_ok(2077, "KG02PR", codigo="REF-9"))
    )
    vistos = []
    monkeypatch.setattr(
        ERPExporter,
        "export",
        lambda self, order, out: vistos.append(order.model_copy(deep=True)) or [],
    )
    cli.main()
    assert [o.items[0].product_code for o in vistos] == ["KG02PR"]
    assert vistos[0].items[0].obs == "ref cliente: REF-9"


def test_detalhe_com_varios_itens_da_planilha(monkeypatch):
    troca = {"seq_previsto": 9, "codigo": "A"}
    check = _check(_ok(10, "A", troca=troca), _ok(11, "B", troca=troca))
    monkeypatch.setattr(product_check, "check_order", lambda o, **kw: check)
    order = _order({"product_code": "A"}, {"product_code": "B"})
    d = preco_trava.decidir_envio(
        "id", order, {}, env=LIGADO, agora=datetime.now(UTC), export_mode="xlsx"
    )
    assert d.detalhe == (
        "Item 1: O importador do Fire trocaria este código pelo produto 9."
        " Mais 1 item(ns) que a planilha não consegue levar ao Fire."
    )
    assert d.block_detail == {
        "produto": [
            {
                "index": 0,
                "estado": "troca_no_fire",
                "motivo": "O importador do Fire trocaria este código pelo produto 9.",
            },
            {
                "index": 1,
                "estado": "troca_no_fire",
                "motivo": "O importador do Fire trocaria este código pelo produto 9.",
            },
        ]
    }


# ── EAN da planilha = EAN do produto decidido (R15b) ────────────────────────


def test_aplicar_decidido_troca_ou_limpa_o_ean_do_cliente():
    order = _order(
        {"product_code": "A", "ean": "1111111111111"}, {"product_code": "B", "ean": "2222"}
    )
    check = _check(_ok(10, "FA", ean="7890000000001"), _ok(20, "FB"))
    depara_apply.aplicar(order, produto_trava.produtos_decididos(check)[0])
    assert [it.ean for it in order.items] == ["7890000000001", None]


def test_aplicar_v1_sem_ean_mantem_o_ean_do_cliente():
    """De-para v1 (fora de `ligado`): comportamento de hoje."""
    order = _order({"product_code": "A", "ean": "1111111111111"})
    v1 = [
        {"item_index": 0, "fire_produto_id": 7, "fire_codigo": "7", "fire_ean": None, "orig": "A"}
    ]
    depara_apply.aplicar(order, v1)
    assert order.items[0].ean == "1111111111111"
    v1[0]["fire_ean"] = ""
    depara_apply.aplicar(order, v1)
    assert order.items[0].ean == "1111111111111"


def test_planilha_leva_o_ean_do_fire_lido_no_check():
    """Casado pelo código; o EAN que o cliente mandou não existe no Fire."""
    produtos = [(2077, "KG02PR", "KIT GRAZI", 1.0, "7891234567890")]
    order = _order({"product_code": "KG02PR", "ean": "7890000000999"})
    check, _ = _check_fire(produtos, order)
    assert check["items"][0]["produto"]["fonte"] == "codprod_altern"
    depara_apply.aplicar(order, produto_trava.produtos_decididos(check)[0])
    (row,) = ERPExporter()._to_erp_rows(order, order.items)
    assert (row.codigo_produto, row.ean) == ("KG02PR", "7891234567890")


# ── Pedido de duas lojas: o split reordena, cada linha leva o seu produto (R15d) ──

LOJA_1 = "11111111000111"
LOJA_2 = "22222222000122"


def _pedido_duas_lojas() -> tuple[Order, dict]:
    order = Order(
        header=OrderHeader(order_number="P2", customer_cnpj="00000000000100", customer_name="ACME"),
        items=[
            OrderItem(
                description=d,
                quantity=1.0,
                unit_price=1.0,
                product_code=f"CLI-{d}",
                delivery_cnpj=loja,
            )
            for d, loja in (("A", LOJA_1), ("B", LOJA_2), ("C", LOJA_1), ("D", LOJA_2))
        ],
    )
    check = _check(
        _ok(10, "FA", codigo="CLI-A", ean="7890000000010"),
        _ok(20, "FB", codigo="CLI-B", ean="7890000000020"),
        _ok(30, "FC", codigo="CLI-C"),
        _ok(40, "FD", codigo="CLI-D", ean="7890000000040"),
    )
    return order, check


ESPERADO = {
    "A": (10, "FA", "7890000000010"),
    "B": (20, "FB", "7890000000020"),
    "C": (30, "FC", None),
    "D": (40, "FD", "7890000000040"),
}


@pytest.mark.parametrize("export_mode", ["xlsx", "db"])
def test_pedido_duas_lojas_cada_linha_leva_o_seu_produto(monkeypatch, tmp_path, export_mode):
    import openpyxl

    from app.exporters import firebird_exporter
    from app.exporters.firebird_exporter import FirebirdExportResult, _to_erp_rows_indexados
    from app.web import server as server_mod

    env = _env_ligado(tmp_path)
    _config(monkeypatch, tmp_path, export_mode)
    _nunca_v1(monkeypatch)
    order, check = _pedido_duas_lojas()
    import_id = _seed(order)
    monkeypatch.setattr(product_check, "check_order", lambda o, **kw: check)
    inseridos = []

    def _export(self, order, *, override_client_id=None, produtos_por_item=None):
        # Mesma indexação do insert real (posição em order.items pela identidade).
        for idx, row in _to_erp_rows_indexados(order):
            inseridos.append((row.descricao, produtos_por_item[idx], row.codigo_produto))
        return FirebirdExportResult("P2", 0, skipped=True, skip_reason="TESTE")

    monkeypatch.setattr(firebird_exporter.FirebirdExporter, "export", _export)
    cfg = {"output_dir": str(tmp_path / "out"), "export_mode": export_mode}
    server_mod._send_one_to_fire(import_id, cfg, request_env=env)

    # Insert: o split por loja reordena (A, C | B, D) e cada linha leva o seu SEQ.
    assert [d for d, _, _ in inseridos] == ["A", "C", "B", "D"]
    assert {d: (seq, cod) for d, seq, cod in inseridos} == {
        d: (seq, cod) for d, (seq, cod, _) in ESPERADO.items()
    }

    planilhas = sorted((tmp_path / "out").glob("*.xlsx"))
    if export_mode == "db":
        assert planilhas == []
        return
    assert len(planilhas) == 2
    linhas = {}
    for p in planilhas:
        ws = openpyxl.load_workbook(p).active
        cab = [c.value for c in ws[1]]
        for r in ws.iter_rows(min_row=2, values_only=True):
            linha = dict(zip(cab, r, strict=True))
            linhas[linha["DESCRICAO"]] = (
                linha["CODIGO_PRODUTO"],
                linha["EAN"] or None,
                linha["CNPJ_LOCAL_ENTREGA"],
                linha["OBS"],
            )
    lojas = {"A": LOJA_1, "B": LOJA_2, "C": LOJA_1, "D": LOJA_2}
    assert set(linhas) == set(ESPERADO)
    for d, (_, cod, ean) in ESPERADO.items():
        codigo, ean_xlsx, loja, obs = linhas[d]
        assert (codigo, ean_xlsx) == (cod, ean)
        assert "".join(ch for ch in str(loja) if ch.isdigit()) == lojas[d]
        assert obs.endswith(f"ref cliente: CLI-{d}")
