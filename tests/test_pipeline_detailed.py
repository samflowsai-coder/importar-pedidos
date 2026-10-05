from pathlib import Path

import openpyxl
import pytest

from app.ingestion.file_loader import LoadedFile
from app.observability.leituras import Diagnostico, cabecalho_candidato

SAMPLES = Path(__file__).parent.parent / "samples"


def _load(filename: str):
    p = SAMPLES / filename
    if not p.exists():
        pytest.skip(f"Sample not found: {filename}")
    return LoadedFile(path=p, extension=p.suffix.lower(), raw=p.read_bytes())


def _resumo(order):
    return order.model_dump(mode="json")


def test_diagnostico_defaults():
    d = Diagnostico()
    assert d.fmt is None and d.parser_usado is None and d.validacao_ok is None
    assert d.extracao == {} and d.parsers_pulados == [] and d.llm == {} and d.avisos == []
    assert d.parser_ms == 0


def test_cabecalho_candidato_puro():
    rows = [[None, "x"], ["IMAGEM", "REF.", "QUAN.", 5], ["A", "B", "C"]]
    assert cabecalho_candidato(rows) == {"linha": 1, "celulas": ["IMAGEM", "REF.", "QUAN."]}
    assert cabecalho_candidato([["a", "1", "b"], [1, 2, 3]]) is None
    assert cabecalho_candidato([]) is None


def test_cabecalho_candidato_ignora_rotulo_curto_e_email():
    rows = [["CT", "NG", "PP", "MM", "GG"], ["IMAGEM", "REF.", "QUAN."], ["a@b.com", "x@y.com", "z@w.com"]]
    assert cabecalho_candidato(rows) == {"linha": 1, "celulas": ["IMAGEM", "REF.", "QUAN."]}


def test_pdf_nao_tem_contagem_de_linhas():
    from app.pipeline import process_detailed

    f = _load("PEDIDO SAMS CLUB.pdf")
    diag = Diagnostico()
    process_detailed(f, diag)
    assert diag.fmt == "pdf" and diag.extracao["linhas"] is None


def test_cabecalho_candidato_limites():
    assert cabecalho_candidato([[None]] * 30 + [["A", "B", "C"]]) is None
    r = cabecalho_candidato([[chr(65 + i) * 3 for i in range(25)]])
    assert len(r["celulas"]) == 20


def test_process_e_detailed_mesma_ordem():
    from app.pipeline import process, process_detailed

    f = _load("Pedido Authentic Feet AF072 QUAN.xlsx")
    assert _resumo(process(f)) == _resumo(process_detailed(f))


def test_af072_quan():
    from app.pipeline import process_detailed

    diag = Diagnostico()
    order = process_detailed(_load("Pedido Authentic Feet AF072 QUAN.xlsx"), diag)
    assert order is not None
    assert diag.parser_usado == "NasmarTemplateParser"
    assert "QUAN." in diag.extracao["cabecalho_candidato"]["celulas"]
    assert diag.extracao["linhas"] > 0 and diag.extracao["abas"]
    assert diag.llm == {"tentado": False, "motivo": "nao_necessario"}
    assert {"nome": "MercadoEletronicoParser", "motivo": "can_parse"} in diag.parsers_pulados
    assert diag.validacao_ok is (diag.avisos == [])
    assert diag.parser_ms >= 0 and diag.fmt


def test_af093_total_kits():
    from app.pipeline import process_detailed

    diag = Diagnostico()
    order = process_detailed(_load("Pedido Authentic Feet AF093 TOTAL KITS.xlsx"), diag)
    assert order is not None
    assert "TOTAL KITS" in diag.extracao["cabecalho_candidato"]["celulas"]


def test_cai_no_generic(tmp_path):
    from app.pipeline import process_detailed

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Codigo", "Descricao", "Quantidade", "Preco"])
    ws.append(["A1", "Tenis azul", 3, 10.5])
    p = tmp_path / "x.xlsx"
    wb.save(p)
    f = LoadedFile(path=p, extension=".xlsx", raw=p.read_bytes())
    diag = Diagnostico()
    order = process_detailed(f, diag)
    assert order is not None
    assert diag.parser_usado == "GenericParser"
    assert diag.parsers_pulados  # os específicos ficaram para trás


def test_motivo_none_e_llm_sem_resultado(monkeypatch):
    import app.pipeline as pl

    class Nulo:
        def can_parse(self, e):
            return True

        def parse(self, e):
            return None

    class LLMNulo:
        def parse(self, e, source_file=None):
            return None

    monkeypatch.setattr(pl, "_parsers", [Nulo()])
    monkeypatch.setattr(pl, "_llm", LLMNulo())
    diag = Diagnostico()
    assert pl.process_detailed(_load("Pedido Authentic Feet AF072 QUAN.xlsx"), diag) is None
    assert diag.parsers_pulados == [{"nome": "Nulo", "motivo": "none"}]
    assert diag.parser_usado is None
    assert diag.llm == {"tentado": True, "motivo": "sem_resultado"}


def test_llm_ok(monkeypatch):
    import app.pipeline as pl
    from app.models.order import Order, OrderHeader, OrderItem

    class LLMOk:
        def parse(self, e, source_file=None):
            return Order(header=OrderHeader(order_number="9"), items=[OrderItem(description="d", quantity=1)])

    monkeypatch.setattr(pl, "_parsers", [])
    monkeypatch.setattr(pl, "_llm", LLMOk())
    diag = Diagnostico()
    assert pl.process_detailed(_load("Pedido Authentic Feet AF072 QUAN.xlsx"), diag) is not None
    assert diag.llm == {"tentado": True, "motivo": "ok"}
    assert diag.parser_usado is None


def test_validador_nao_bloqueia(monkeypatch):
    import app.pipeline as pl

    monkeypatch.setattr(pl._validator, "validate_detail", lambda o: ["ruim"])
    diag = Diagnostico()
    assert pl.process_detailed(_load("Pedido Authentic Feet AF072 QUAN.xlsx"), diag) is not None
    assert diag.avisos == ["ruim"] and diag.validacao_ok is False


def test_sheets_xlsx(tmp_path):
    from app.extractors.xls_extractor import XLSExtractor

    wb = openpyxl.Workbook()
    wb.active.title = "A"
    wb.active.append(["x"])
    wb.active.append([None])
    wb.active.append(["y"])
    wb.create_sheet("B").append(["z"])
    p = tmp_path / "s.xlsx"
    wb.save(p)
    out = XLSExtractor().extract(LoadedFile(path=p, extension=".xlsx", raw=p.read_bytes()))
    assert out["sheets"] == [{"name": "A", "rows": 2}, {"name": "B", "rows": 1}]
    assert len(out["rows"]) == 3


def test_sheets_xls(monkeypatch):
    import app.extractors.xls_extractor as mod

    class Sheet:
        def __init__(self, name, data):
            self.name, self._d = name, data
            self.nrows = len(data)

        def row_values(self, i):
            return self._d[i]

    class WB:
        def sheets(self):
            return [Sheet("P1", [["a"], ["b"]]), Sheet("P2", [["c"]])]

    monkeypatch.setattr(mod.xlrd, "open_workbook", lambda file_contents: WB())
    out = mod.XLSExtractor().extract(LoadedFile(path=Path("a.xls"), extension=".xls", raw=b""))
    assert out["sheets"] == [{"name": "P1", "rows": 2}, {"name": "P2", "rows": 1}]
